# -*- coding: utf-8 -*-
"""
고객안내문(FSN) + FCO 통합 Dashboard

SharePoint(OneDrive)에 있는 고객안내문 전달 현황 엑셀과 FCO Action Plan 엑셀을 받아서
브라우저(Edge)에서 보는 HTML 대시보드로 정리합니다.

- 더블클릭(고객안내문.exe) → 오늘 받은 파일이 없으면 Edge로 자동 다운로드 → 대시보드 열기
- 데이터/로그 폴더: ~/FSN_Dashboard/
- 개발용: python main.py --build 고객안내문.xlsx --fco MR=MR.xlsx --out test.html --no-open
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import time
import traceback
import webbrowser
import zipfile
from collections import Counter, OrderedDict
from pathlib import Path

APP_TITLE = "고객안내문 · FCO 통합 Dashboard"
APP_VERSION = "1.0.0"

HOME = Path.home()
APP_DIR = HOME / "FSN_Dashboard"
DATA_DIR = APP_DIR / "data"
DOWNLOAD_DIR = APP_DIR / "downloads"
PROFILE_DIR = APP_DIR / "browser_profile"
LOG_PATH = APP_DIR / "debug_log.txt"
CONFIG_PATH = APP_DIR / "config.json"
DASHBOARD_PATH = APP_DIR / "FSN_FCO_Dashboard.html"
FCO_APP_DIR = HOME / "FCO_Action_Plan"  # 기존 FCO 프로그램의 캐시 폴더 (읽기만 함)
TEMPLATE_NAME = "dashboard_template.html"

# 고객안내문 전달 현황 파일
FSN_SOURCE = {
    "name": "고객안내문",
    "url": "https://docs.philips.com/:x:/g/personal/donghak_lee_philips_com/"
           "IQD_cfiaUjaWTbbrDq6Ki7PrAVaUk2adIR0GtUH-QSYvpDo?e=PzbYYK",
    "cache": "FSN_customer_notice.xlsx",
}

# FCO Action Plan 파일 (KRHSFCOstatus 사이트)
# url이 비어 있으면 기존 FCO 프로그램(~/FCO_Action_Plan)이 받아 둔 파일을 사용합니다.
# 주소는 config.json의 "fco_urls"로도 넣을 수 있습니다.
FCO_SOURCES = OrderedDict([
    ("MR", {"url": "", "cache": "MR_FCO_Action_Plan.xlsx", "pattern": r"(?<![A-Za-z])MR(?![a-z])"}),
    ("XR", {"url": "", "cache": "XR_Open_FCO_list.xlsx", "pattern": r"(?<![A-Za-z])XR(?![a-z])"}),
    ("CT/AMI", {"url": "", "cache": "CT_AMI_FCO_Action_Plan.xlsx", "pattern": r"(?<![A-Za-z])CT(?![a-z])"}),
])

EXCEL_EXTS = (".xlsx", ".xlsm")


# ---------------------------------------------------------------------------
# 로그 / 공용 도우미
# ---------------------------------------------------------------------------

def ensure_dirs():
    for d in (APP_DIR, DATA_DIR, DOWNLOAD_DIR):
        d.mkdir(parents=True, exist_ok=True)


def log(msg: str):
    try:
        APP_DIR.mkdir(parents=True, exist_ok=True)
        if LOG_PATH.exists() and LOG_PATH.stat().st_size > 2_000_000:
            old = LOG_PATH.with_suffix(".old.txt")
            try:
                if old.exists():
                    old.unlink()
                LOG_PATH.rename(old)
            except OSError:
                pass
        with open(LOG_PATH, "a", encoding="utf-8") as f:
            f.write(f"[{dt.datetime.now():%Y-%m-%d %H:%M:%S}] {msg}\n")
    except Exception:
        pass


def log_exc(prefix: str):
    log(prefix + "\n" + traceback.format_exc())


def resource_path(name: str) -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / name


def load_config() -> dict:
    """~/FSN_Dashboard/config.json (선택). 없으면 빈 설정."""
    try:
        if CONFIG_PATH.exists():
            with open(CONFIG_PATH, encoding="utf-8-sig") as f:
                cfg = json.load(f)
            if isinstance(cfg, dict):
                return cfg
    except Exception:
        log_exc("config.json 읽기 실패")
    return {}


def safe_filename(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", s).strip("_") or "file"


_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean_text(v) -> str:
    """셀 값을 화면용 문자열로 변환 (제어문자 제거)."""
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:g}"
    s = _CTRL_RE.sub("", str(v))
    return s.strip()


def norm_header(v) -> str:
    """헤더 비교용: 소문자, 공백/구두점 제거 (&는 유지)."""
    if v is None:
        return ""
    s = clean_text(v).lower().replace("＆", "&")
    return re.sub(r"[\s_\-.:()\[\]{}#/\\·,'\"’`~※*]+", "", s)


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        if 20000 < float(v) < 80000:  # 엑셀 날짜 일련번호
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
        return None
    s = clean_text(v)
    if not s:
        return None
    m = re.search(r"(20\d{2}|19\d{2})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})", s)
    if m:
        try:
            return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
    m = re.fullmatch(r"(\d{2})[.\-/](\d{1,2})[.\-/](\d{1,2})\.?", s)
    if m:
        try:
            return dt.date(2000 + int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
    if not re.search(r"\d", s):
        return None
    try:
        from dateutil import parser as dparser
        return dparser.parse(s, fuzzy=False).date()
    except Exception:
        return None


def split_people(s: str) -> list:
    if not s:
        return []
    parts = re.split(r"[,/;\n·&+]|\s{2,}", s)
    out = []
    for p in parts:
        p = p.strip(" ()[]")
        if p and p not in out and p not in ("-", "?"):
            out.append(p)
    return out


_CORP_RE = re.compile(r"\((주|재|의|학|사|복|유|합|사단|재단)\)|㈜|㈔|재단법인|의료법인|학교법인|사회복지법인|사단법인|주식회사")


def hospital_key(name: str) -> str:
    s = clean_text(name).lower()
    s = _CORP_RE.sub("", s)
    s = re.sub(r"[\s\-_.,·()\[\]/'\"]+", "", s)
    return s


def sn_key(sn: str) -> str:
    return re.sub(r"[^0-9A-Za-z]", "", clean_text(sn)).upper()


# ---------------------------------------------------------------------------
# 엑셀 읽기: 시트/헤더/컬럼 자동 인식
# ---------------------------------------------------------------------------

def _rule(exact=(), contains=(), exclude=()):
    return {
        "exact": [norm_header(a) for a in exact],
        "contains": [norm_header(a) for a in contains],
        "exclude": [norm_header(a) for a in exclude],
    }


# 순서가 우선순위입니다 (앞의 필드가 먼저 컬럼을 차지).
FSN_FIELDS = [
    ("hospital", _rule(
        exact=["국문병원명", "병원명(국문)", "국문 병원명", "병원명 국문", "고객명(국문)", "국문고객명", "고객명 국문",
               "병원명", "병원", "병원 이름", "고객명", "고객사명", "고객사", "고객", "의료기관명", "의료기관",
               "기관명", "기관", "거래처명", "거래처", "Account Name", "Account", "Customer Name",
               "Customer", "Hospital", "Hospital Name", "Site Name", "Site"],
        contains=["병원", "고객사", "고객명", "기관명", "거래처", "account", "customer", "hospital"],
        exclude=["담당", "코드", "code", "번호", "주소", "address", "연락", "전화", "phone", "안내", "notice",
                 "유형", "type", "id", "수량", "개수", "구분", "종별", "등급", "영문", "english"])),
    # 병원명 칸이 비었을 때 쓸 두 번째 국문 이름 (예: 국문병원명이 비면 고객명(국문))
    ("hospital2", _rule(
        exact=["고객명(국문)", "국문고객명", "고객명 국문", "국문병원명", "병원명(국문)", "고객명", "고객사명"],
        contains=["국문"], exclude=["영문", "코드", "주소"])),
    # 국문 이름이 모두 비었을 때 마지막 대안
    ("hospital_en", _rule(
        exact=["영문병원명", "영문 병원명", "병원명(영문)", "고객명(영문)", "영문고객명", "Account Name",
               "Customer Name", "Hospital Name"],
        contains=["영문"], exclude=["코드", "주소"])),
    ("method", _rule(
        exact=["전달방법", "전달 방식", "전달수단", "전달경로", "발송방법", "발송방식", "송부방법", "송부방식",
               "통보방법", "방법", "Delivery Method", "Method", "Delivery Type"],
        contains=["방법", "방식", "수단", "method"])),
    ("status", _rule(
        exact=["전달여부", "전달 유무", "전달현황", "전달상태", "전달결과", "발송여부", "발송현황", "발송상태",
               "송부여부", "완료여부", "진행상태", "진행현황", "진행상황", "상태", "Status", "결과", "전달",
               "발송", "완료", "Delivered", "Delivery Status", "Delivery"],
        contains=["여부", "상태", "status", "현황", "유무", "결과"],
        exclude=["회신", "서명", "ack", "확인서", "답신", "응답", "c&r"])),
    ("date", _rule(
        exact=["전달일", "전달일자", "전달날짜", "전달 일시", "발송일", "발송일자", "송부일", "송부일자", "완료일",
               "완료일자", "수령일", "통보일", "Delivery Date", "Sent Date", "Date", "일자", "날짜"],
        contains=["전달일", "발송일", "송부일", "완료일", "수령일", "통보일", "deliverydate", "sentdate"],
        exclude=["회신", "예정", "due", "발행"])),
    ("reply", _rule(
        exact=["회신", "회신여부", "회신서", "회신일", "회신 유무", "서명", "서명여부", "확인서", "Acknowledgement",
               "Ack", "ACK Form"],
        contains=["회신", "서명", "확인서", "답신", "ack", "acknowledg"])),
    ("site", _rule(exact=["Site", "Site Name", "사이트", "설치장소", "설치 위치", "부서", "Location"])),
    ("cr", _rule(
        exact=["C&R", "CR", "C&R No", "C&R 번호", "CR No", "C&R 구분", "C&R ID", "C&R Number", "C&R명"],
        contains=["c&r"])),
    ("bu", _rule(
        exact=["BU", "Business Unit", "사업부", "BG", "Business Group", "Modality", "모달리티", "제품군", "BU 구분",
               "분야"],
        contains=["bu구분", "사업부", "modality", "모달리티", "businessunit"])),
    ("fsn", _rule(
        exact=["FSN", "FSN No", "FSN 번호", "FSN Number", "FSN ID", "FSN명", "FSN 코드", "안내문 번호",
               "고객안내문 번호", "고객안내문", "고객안내문명", "안내문", "안내문명", "FCA", "FCA No",
               "Field Action", "Field Action No"],
        contains=["fsn", "안내문", "fieldaction"],
        exclude=["여부", "상태", "현황", "일자", "날짜", "방법", "방식", "전달", "발송", "회신", "status", "date",
                 "method", "제목", "title"])),
    ("fsn_title", _rule(
        exact=["제목", "Title", "FSN 제목", "FSN Title", "안내문 제목", "내용", "Subject", "FSN 내용"],
        contains=["제목", "title", "subject"])),
    ("fco", _rule(exact=["FCO", "FCO No", "FCO 번호", "관련 FCO", "FCO Number"],
                  contains=["fco"], exclude=["제목", "title", "여부", "상태"])),
    ("fse", _rule(
        exact=["FSE", "담당 FSE", "FSE명", "FSE 담당", "FSE 이름", "엔지니어", "담당 엔지니어", "Engineer", "CSE",
               "서비스 담당", "담당자", "담당"],
        contains=["fse", "엔지니어", "engineer"])),
    ("sales", _rule(
        exact=["Sales", "영업", "영업담당", "영업 담당자", "세일즈", "KAM", "AM", "Account Manager", "Sales Rep",
               "Sales 담당"],
        contains=["sales", "영업", "세일즈"])),
    ("sn", _rule(
        exact=["SN", "S/N", "Serial", "Serial No", "Serial Number", "시리얼", "시리얼번호", "장비번호", "Equipment No"],
        contains=["serial", "시리얼", "일련번호"])),
    ("system", _rule(
        exact=["System Code Description", "System", "System Code", "장비", "장비명", "모델", "모델명", "Model",
               "Product", "제품", "제품명", "Equipment", "장비모델", "시스템"],
        contains=["장비명", "모델", "model", "systemcode", "product", "제품명"])),
    ("region", _rule(exact=["Region", "지역", "권역", "지사", "지점", "Area", "팀", "Team"],
                     contains=["region", "지역", "권역"])),
    ("remark", _rule(exact=["비고", "Remark", "Remarks", "메모", "Note", "Notes", "Comment", "Comments", "특이사항",
                            "코멘트"],
                     contains=["비고", "remark", "특이", "comment", "메모"])),
]

# 전달방법이 컬럼 하나가 아니라 "등기 | 직접전달 | 이메일" 처럼 따로 있는 경우
WIDE_METHOD_HEADERS = {
    norm_header(k): v for k, v in {
        "등기": "등기", "등기우편": "등기", "우편": "등기", "등기발송": "등기",
        "직접전달": "직접전달", "직접": "직접전달", "방문": "직접전달", "방문전달": "직접전달", "직접방문": "직접전달",
        "이메일": "이메일", "Email": "이메일", "E-mail": "이메일", "메일": "이메일",
    }.items()
}

FCO_FIELDS = [
    ("fco", _rule(exact=["FCO", "FCO No", "FCO 번호", "FCO Number"])),
    ("due", _rule(exact=["Due Date", "Due", "마감일", "기한", "Due Dt"], contains=["due"])),
    ("title", _rule(exact=["FCO Title", "Title", "제목"], contains=["title"])),
    ("sysdesc", _rule(exact=["System Code Description", "System Description"])),
    ("syscode", _rule(exact=["System Code"])),
    ("sn", _rule(exact=["SN", "S/N", "Serial", "Serial No"])),
    ("account", _rule(exact=["Account", "Account Name", "병원", "병원명", "고객", "고객명"])),
    ("site", _rule(exact=["Site", "Site Name"])),
    ("case", _rule(exact=["Case", "Case No"])),
    ("wo", _rule(exact=["WO", "WO No", "Work Order"])),
    ("region", _rule(exact=["Region", "지역"])),
    ("fse", _rule(exact=["FSE", "담당", "담당자", "FSE명"])),
    ("ap", _rule(exact=["Action Plan"], contains=["actionplan"])),
    ("status", _rule(exact=["Status", "상태"])),
    ("rel", _rule(exact=["고객안내문 관련", "고객안내문"], contains=["안내문"])),
    ("remark", _rule(exact=["Remark", "Remarks", "비고"])),
]


def match_columns(header_cells, spec, overrides=None):
    """header_cells: [(col_idx, text)] → {field: col_idx}"""
    normed = [(c, norm_header(t)) for c, t in header_cells]
    normed = [(c, h) for c, h in normed if h]
    assigned, used = {}, set()
    for field, want in (overrides or {}).items():
        w = norm_header(want)
        for c, h in normed:
            if c not in used and h == w:
                assigned[field] = c
                used.add(c)
                break
    for field, rule in spec:
        if field in assigned:
            continue
        for alias in rule["exact"]:
            hit = next((c for c, h in normed if c not in used and h == alias), None)
            if hit is not None:
                assigned[field] = hit
                used.add(hit)
                break
    for field, rule in spec:
        if field in assigned or not rule["contains"]:
            continue
        for c, h in normed:
            if c in used:
                continue
            if any(w in h for w in rule["contains"]) and not any(x in h for x in rule["exclude"]):
                assigned[field] = c
                used.add(c)
                break
    return assigned


def read_sheet_grids(ws):
    """(모든 병합 채움 grid, 세로 병합만 채운 grid) 반환. grid[r][c] 0-based."""
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    width = max((len(r) for r in raw), default=0)
    for r in raw:
        if len(r) < width:
            r.extend([None] * (width - len(r)))
    grid_all = [list(r) for r in raw]
    grid_vert = [list(r) for r in raw]
    try:
        ranges = list(ws.merged_cells.ranges)
    except Exception:
        ranges = []
    for rng in ranges:
        r0, c0, r1, c1 = rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1
        if r0 >= len(raw) or c0 >= width:
            continue
        v = raw[r0][c0]
        for r in range(r0, min(r1, len(raw) - 1) + 1):
            for c in range(c0, min(c1, width - 1) + 1):
                if grid_all[r][c] is None:
                    grid_all[r][c] = v
                if c0 == c1 and grid_vert[r][c] is None:
                    grid_vert[r][c] = v
    return grid_all, grid_vert


def detect_header(grid, spec, required, overrides=None, scan_rows=25):
    """가장 많은 필드가 매칭되는 헤더 행을 찾음. 두 줄 헤더도 시도.
    반환: dict(header_row(0-based 마지막 헤더행), headers=[(c,text)], mapping, score) 또는 None"""
    best = None
    n = min(scan_rows, len(grid))
    for r in range(n):
        candidates = [(r, [(c, clean_text(v)) for c, v in enumerate(grid[r])])]
        distinct = {clean_text(v) for v in grid[r] if clean_text(v)}
        # 제목 행(값이 하나뿐인 행)은 두 줄 헤더 조합에서 제외
        if r + 1 < len(grid) and len(distinct) > 1:
            comb = []
            for c in range(len(grid[r])):
                top = clean_text(grid[r][c])
                bottom = clean_text(grid[r + 1][c])
                comb.append((c, bottom or top))
            candidates.append((r + 1, comb))
        for last_row, cells in candidates:
            mapping = match_columns(cells, spec, overrides)
            if not all(any(f in mapping for f in group) for group in required):
                continue
            wide = {c: WIDE_METHOD_HEADERS[norm_header(t)] for c, t in cells
                    if c not in mapping.values() and norm_header(t) in WIDE_METHOD_HEADERS}
            # 등기/직접전달/이메일 칸이 2개 이상이면 "전달방법 + 칸 구조" 2점
            score = len(mapping) + (2 if len(wide) >= 2 else 0)
            # 두 줄 헤더는 더 많이 맞을 때만 채택
            if last_row != r and best is not None and best["first_row"] == r and score <= best["score"]:
                continue
            if best is None or score > best["score"]:
                best = {"first_row": r, "header_row": last_row, "headers": cells, "mapping": mapping,
                        "wide": wide, "score": score}
    return best


SKIP_HOSPITAL_WORDS = {"합계", "총계", "소계", "계", "total", "sum", "subtotal", "grandtotal"}


# ---------------------------------------------------------------------------
# 고객안내문(FSN) 파싱
# ---------------------------------------------------------------------------

EXCLUDE_STATUS = ["해당없음", "대상아님", "대상외", "제외", "n/a", "철거", "폐기", "불필요", "취소"]
DONE_EXACT = {"o", "y", "yes", "true", "1", "v", "✓", "✔", "○", "◯", "●", "◎", "ok", "done", "완료", "전달", "발송",
              "송부", "수령", "전달완료", "발송완료", "송부완료", "delivered", "sent", "complete", "completed",
              "closed", "통보", "통보완료"}
TODO_EXACT = {"x", "n", "no", "false", "0", "×", "✗", "✘", "-", "미", "대기", "예정", "진행", "진행중", "보류",
              "pending", "open", "todo", "notyet", "미정", "na"}
TODO_SUB = ["미전달", "미발송", "미송부", "미완료", "미수령", "미진행", "미통보", "안됨", "못함", "대기", "예정", "진행중",
            "보류", "반송", "거부", "거절", "불가", "pending", "notyet", "notdelivered", "waiting", "hold"]
DONE_SUB = ["완료", "전달", "발송", "송부", "수령", "통보", "등기", "우편", "직접", "방문", "이메일", "메일", "email",
            "fax", "팩스", "deliver", "sent", "complete", "done"]


def classify_status(v):
    """→ 'done' | 'todo' | 'excluded' | 'unknown' | '' (빈값)"""
    if v is None:
        return ""
    if isinstance(v, (dt.date, dt.datetime)):
        return "done"
    s = norm_header(v)
    if not s:
        return ""
    if any((w == s) if len(w) <= 3 else (w in s) for w in (norm_header(x) for x in EXCLUDE_STATUS)):
        return "excluded"
    if s in TODO_EXACT:
        return "todo"
    if s in DONE_EXACT:
        return "done"
    if s.startswith("미") or any(w in s for w in TODO_SUB):
        return "todo"
    if any(w in s for w in DONE_SUB):
        return "done"
    if to_date(v):
        return "done"
    return "unknown"


KNOWN_METHODS = ("등기", "직접전달", "이메일", "팩스")


def canon_method(v) -> str:
    s = norm_header(v)
    if not s:
        return ""
    found = []
    if any(w in s for w in ("등기", "우편", "우체국", "택배", "registered", "post", "courier")):
        found.append("등기")
    if any(w in s for w in ("직접", "방문", "대면", "인편", "hand", "visit", "direct")):
        found.append("직접전달")
    s_mail = s.replace("registeredmail", "")
    if any(w in s_mail for w in ("이메일", "메일", "email", "mail")):
        found.append("이메일")
    if any(w in s for w in ("팩스", "fax")):
        found.append("팩스")
    if found:
        return "+".join(found)
    if s in TODO_EXACT or s.startswith("미"):
        return ""
    return clean_text(v)[:30]


def norm_bu(v) -> str:
    s = clean_text(v)
    if s.isascii() and len(s) <= 8:
        return s.upper()
    return s


def _truthy_mark(v) -> bool:
    if v is None:
        return False
    if isinstance(v, (dt.date, dt.datetime)):
        return True
    s = norm_header(v)
    return bool(s) and s not in TODO_EXACT and not s.startswith("미")


HOSP_FIELDS = ("hospital", "hospital2", "hospital_en")
PLACEHOLDER = {"", "-", "--", "n", "na", "x", "0", "?", "없음", "미정", "tbd"}
# 이 단어가 들어간 시트는 읽지 않음 (config.json "skip_sheets"로 바꿀 수 있음)
SKIP_SHEET_WORDS = ["삭제", "closed"]


def row_hospital(row, m) -> str:
    """국문병원명 → 고객명(국문) → 영문병원명 순으로 처음 채워진 값"""
    for f in HOSP_FIELDS:
        c = m.get(f)
        if c is None or c >= len(row):
            continue
        v = clean_text(row[c])
        if norm_header(v) not in PLACEHOLDER:
            return v
    return ""


def parse_fsn(path: Path, overrides=None, skip_sheets=None) -> dict:
    import openpyxl
    res = {"ok": False, "error": "", "rows": [], "sheets": [], "status_values": [], "method_values": [],
           "excluded": 0, "skipped": 0, "unknown_status": 0, "warnings": [], "fsn_from_sheet": False,
           "fields": {}}
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        res["error"] = f"엑셀 파일을 열 수 없습니다: {e}"
        log_exc("FSN 파일 열기 실패")
        return res

    candidates = []
    skip_words = [norm_header(w) for w in (skip_sheets if skip_sheets is not None else SKIP_SHEET_WORDS)]
    for ws in wb.worksheets:
        info = {"name": ws.title, "used": False, "reason": "", "rows": 0}
        res["sheets"].append(info)
        if getattr(ws, "sheet_state", "visible") != "visible":
            info["reason"] = "숨김 시트"
            continue
        if any(w and w in norm_header(ws.title) for w in skip_words):
            info["reason"] = "시트 이름으로 제외 (삭제예정/Closed 등)"
            continue
        try:
            grid_all, grid_vert = read_sheet_grids(ws)
        except Exception as e:
            info["reason"] = f"읽기 실패: {e}"
            continue
        hdr = detect_header(grid_all, FSN_FIELDS, required=[HOSP_FIELDS], overrides=overrides)
        if not hdr:
            info["reason"] = "병원 컬럼을 찾지 못함"
            continue
        m = hdr["mapping"]
        if not (("status" in m) or ("method" in m) or ("date" in m) or hdr["wide"]):
            info["reason"] = "전달여부/전달방법/전달일 컬럼을 찾지 못함"
            info["header_row"] = hdr["header_row"] + 1
            info["mapping"] = {f: hdr["headers"][c][1] for f, c in m.items()}
            continue
        data_rows = sum(1 for row in grid_vert[hdr["header_row"] + 1:] if row_hospital(row, m))
        info.update({"header_row": hdr["header_row"] + 1, "score": hdr["score"], "data_rows": data_rows,
                     "mapping": {f: hdr["headers"][c][1] for f, c in m.items()},
                     "wide": {hdr["headers"][c][1]: v for c, v in hdr["wide"].items()}})
        candidates.append((ws, info, hdr, grid_vert))

    if not candidates:
        res["error"] = "고객안내문 데이터 시트를 찾지 못했습니다. (병원명 + 전달여부/전달방법 컬럼이 필요합니다)"
        return res

    best = max(candidates, key=lambda x: (x[2]["score"], x[1]["data_rows"]))
    best_fields = set(best[2]["mapping"].keys())
    # 같은 구조(핵심 컬럼이 같은) 시트만 합침 → 요약/피벗 시트는 제외
    core = best_fields & {"status", "method", "date", "bu", "cr", "fsn"}
    chosen = [c for c in candidates
              if core <= set(c[2]["mapping"].keys()) and (bool(c[2]["wide"]) == bool(best[2]["wide"]))]
    # 시트가 여러 개일 때: BU 컬럼이 없으면 시트 이름 = BU (예: MR / CT / IGT DXR US HPM 시트)
    #                     BU는 있고 FSN/C&R 컬럼이 없으면 시트 이름 = FSN
    bu_from_sheet = len(chosen) > 1 and "bu" not in best_fields
    if len(chosen) > 1 and not bu_from_sheet and not ({"fsn", "cr"} & best_fields):
        res["fsn_from_sheet"] = True

    status_counter, method_counter = Counter(), Counter()
    seen = set()
    for ws, info, hdr, grid in chosen:
        info["used"] = True
        m, wide = hdr["mapping"], hdr["wide"]
        cnt = 0

        def g(row, field):
            c = m.get(field)
            if c is None or c >= len(row):
                return None
            return row[c]

        for ri in range(hdr["header_row"] + 1, len(grid)):
            row = grid[ri]
            hosp = row_hospital(row, m)
            if not hosp:
                if any(clean_text(v) for v in row):
                    res["skipped"] += 1
                continue
            if norm_header(hosp) in SKIP_HOSPITAL_WORDS or norm_header(hosp).startswith("합계"):
                continue
            # 반복 헤더 행 건너뛰기
            if norm_header(hosp) in {norm_header(hdr["headers"][m[f]][1]) for f in HOSP_FIELDS if f in m}:
                continue

            status_raw = g(row, "status")
            date_v = to_date(g(row, "date"))
            method_raw = clean_text(g(row, "method"))
            method = canon_method(method_raw)
            # 등기/직접전달/이메일이 컬럼으로 나뉜 경우: 표시(O, 날짜 등)가 있는 컬럼이 전달방법
            wide_methods = [mv for c, mv in wide.items() if c < len(row) and _truthy_mark(row[c])]
            if not method and wide_methods:
                method = "+".join(dict.fromkeys(wide_methods))
            if not date_v:
                for c in wide:
                    if c < len(row) and isinstance(row[c], (dt.date, dt.datetime)):
                        date_v = to_date(row[c])
                        break

            if "status" in m:
                state = classify_status(status_raw)
                status_counter[(clean_text(status_raw) or "(빈칸)", state or "empty")] += 1
                if state == "excluded":
                    res["excluded"] += 1
                    continue
                if state == "unknown":
                    res["unknown_status"] += 1
                if state in ("", "unknown"):
                    state = "done" if (date_v or wide_methods) else "todo"
                # 전달여부 칸에 "등기", "이메일" 처럼 방법이 적힌 경우
                if not method and state == "done":
                    cm = canon_method(status_raw)
                    if cm and all(p in KNOWN_METHODS for p in cm.split("+")):
                        method = cm
            elif "date" in m:
                state = "done" if (date_v or wide_methods) else "todo"
            else:
                state = "done" if (method or wide_methods) else "todo"

            if method_raw or method:
                method_counter[(method_raw or "+".join(wide_methods), method or "(미지정)")] += 1

            fsn = clean_text(g(row, "fsn"))
            if not fsn and res["fsn_from_sheet"]:
                fsn = ws.title
            rec = {
                "k": "", "h": hosp,
                "site": clean_text(g(row, "site")),
                "bu": norm_bu(g(row, "bu")) or (ws.title.strip() if bu_from_sheet else ""),
                "hen": clean_text(g(row, "hospital_en")) if row_hospital(row, m) != clean_text(
                    g(row, "hospital_en")) else "",
                "cr": clean_text(g(row, "cr")),
                "fsn": fsn,
                "ft": clean_text(g(row, "fsn_title")),
                "fco": clean_text(g(row, "fco")),
                "m": method,
                "mr": method_raw,
                "s": 1 if state == "done" else 0,
                "sr": clean_text(status_raw),
                "d": date_v.isoformat() if date_v else "",
                "rp": clean_text(g(row, "reply")),
                "fse": clean_text(g(row, "fse")),
                "sales": clean_text(g(row, "sales")),
                "sn": clean_text(g(row, "sn")),
                "sys": clean_text(g(row, "system")),
                "rg": clean_text(g(row, "region")),
                "rm": clean_text(g(row, "remark"))[:300],
                "sheet": ws.title,
                "row": ri + 1,
            }
            dedupe_key = tuple(v for kk, v in rec.items() if kk not in ("sheet", "row"))
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)
            res["rows"].append(rec)
            cnt += 1
        info["rows"] = cnt

    res["fields"] = best[1]["mapping"]
    res["wide"] = best[1].get("wide", {})
    res["status_values"] = [{"value": v, "state": s, "count": n}
                            for (v, s), n in status_counter.most_common(40)]
    res["method_values"] = [{"value": v, "canon": cn, "count": n}
                            for (v, cn), n in method_counter.most_common(30)]
    if "status" not in best[2]["mapping"]:
        if "date" in best[2]["mapping"]:
            res["warnings"].append("전달여부 컬럼이 없어 '전달일'이 입력된 행을 완료로 계산했습니다.")
        else:
            res["warnings"].append("전달여부 컬럼이 없어 '전달방법'이 입력된 행을 완료로 계산했습니다.")
    if res["unknown_status"]:
        res["warnings"].append(f"전달여부 값을 해석하지 못한 행 {res['unknown_status']}건은 미전달로 계산했습니다. "
                               "(데이터 정보 탭의 '전달여부 값' 표 참고)")
    if res["fsn_from_sheet"]:
        res["warnings"].append("FSN 번호 컬럼이 없어 시트 이름을 FSN 구분으로 사용했습니다.")
    if bu_from_sheet:
        res["warnings"].append("BU 컬럼이 없어 시트 이름(" + ", ".join(c[1]["name"] for c in chosen) +
                               ")을 BU로 사용했습니다.")
    res["ok"] = True
    log(f"FSN 파싱: {len(res['rows'])}행, 시트={[c[1]['name'] for c in chosen]}, 매핑={res['fields']}")
    return res


# ---------------------------------------------------------------------------
# FCO 파싱
# ---------------------------------------------------------------------------

def due_band(due, today):
    from dateutil.relativedelta import relativedelta
    if not due:
        return "n"
    if due < today + relativedelta(months=6):
        return "r"
    if due < today + relativedelta(years=1):
        return "o"
    return "g"


def parse_fco(path: Path, source: str, today: dt.date) -> dict:
    import openpyxl
    res = {"ok": False, "error": "", "rows": [], "sheet": "", "header_row": 0, "fields": {}, "closed": 0}
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        res["error"] = f"엑셀 파일을 열 수 없습니다: {e}"
        log_exc(f"FCO({source}) 파일 열기 실패")
        return res
    best = None
    try:
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            grid = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                grid.append(list(r))
            hdr = detect_header(grid, FCO_FIELDS, required=[("fco",), ("due", "fse", "account")])
            if hdr and (best is None or hdr["score"] > best[1]["score"]):
                best = (ws.title, hdr, grid)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    if not best:
        res["error"] = "FCO 데이터 시트를 찾지 못했습니다. (FCO, Due Date 컬럼 필요)"
        return res
    title, hdr, grid = best
    m = hdr["mapping"]
    res["sheet"], res["header_row"] = title, hdr["header_row"] + 1
    res["fields"] = {f: hdr["headers"][c][1] for f, c in m.items()}

    def g(row, field):
        c = m.get(field)
        if c is None or c >= len(row):
            return None
        return row[c]

    for ri in range(hdr["header_row"] + 1, len(grid)):
        row = grid[ri]
        fco = clean_text(g(row, "fco"))
        if not fco:
            continue
        if norm_header(fco) == "fco":
            continue
        status = clean_text(g(row, "status"))
        if status.strip().lower() == "closed":
            res["closed"] += 1
            continue
        due = to_date(g(row, "due"))
        account = clean_text(g(row, "account"))
        site = clean_text(g(row, "site"))
        res["rows"].append({
            "k": "", "h": account or site, "site": site, "src": source, "fco": fco,
            "due": due.isoformat() if due else "", "dd": (due - today).days if due else None,
            "b": due_band(due, today),
            "t": clean_text(g(row, "title")),
            "sn": clean_text(g(row, "sn")),
            "sys": clean_text(g(row, "sysdesc")) or clean_text(g(row, "syscode")),
            "fse": clean_text(g(row, "fse")),
            "ap": clean_text(g(row, "ap"))[:300],
            "st": status,
            "rel": clean_text(g(row, "rel")),
            "wo": clean_text(g(row, "wo")),
            "case": clean_text(g(row, "case")),
            "rg": clean_text(g(row, "region")),
            "rm": clean_text(g(row, "remark"))[:300],
            "row": ri + 1,
        })
    res["ok"] = True
    log(f"FCO({source}) 파싱: {len(res['rows'])}건 (Closed 제외 {res['closed']}), 시트={title}")
    return res


# ---------------------------------------------------------------------------
# 대시보드 데이터 만들기 + HTML 생성
# ---------------------------------------------------------------------------

def build_dashboard_data(fsn_info: dict, fco_infos: list, today: dt.date) -> dict:
    """fsn_info: {"parse": parse_fsn 결과, "file": {...}} / fco_infos: [{"source","parse","file"}]"""
    fsn_parse = (fsn_info or {}).get("parse") or {}
    fsn_rows = fsn_parse.get("rows", []) if fsn_parse.get("ok") else []
    fco_rows = []
    for info in fco_infos:
        p = info.get("parse") or {}
        if p.get("ok"):
            fco_rows.extend(p["rows"])

    # 병원 키: FCO 장비 S/N → Account 매핑을 이용해 이름이 달라도 같은 병원으로 묶음
    sn_to_key = {}
    for r in fco_rows:
        r["k"] = hospital_key(r["h"]) or ("fco:" + r["fco"])
        if r["sn"]:
            sn_to_key.setdefault(sn_key(r["sn"]), r["k"])
    linked_by_sn = 0
    fco_keys = {r["k"] for r in fco_rows}
    for r in fsn_rows:
        k = ""
        if r["sn"] and sn_key(r["sn"]) in sn_to_key:
            k = sn_to_key[sn_key(r["sn"])]
            if k != hospital_key(r["h"]):
                linked_by_sn += 1
        if not k and hospital_key(r["h"]) not in fco_keys and r.get("hen") and hospital_key(r["hen"]) in fco_keys:
            k = hospital_key(r["hen"])  # FCO Account가 영문 이름인 경우
        r["k"] = k or hospital_key(r["h"]) or ("fsn:" + r["h"])

    hospitals = {}
    for r in fsn_rows:
        h = hospitals.setdefault(r["k"], {"name": r["h"], "pp": [], "names": []})
        for nm in (r["h"], r.get("hen")):
            if nm and nm not in h["names"]:
                h["names"].append(nm)
        pp = split_people(r["fse"]) + split_people(r["sales"])
        r["pp"] = pp
        for p in pp:
            if p not in h["pp"]:
                h["pp"].append(p)
    for r in fco_rows:
        h = hospitals.setdefault(r["k"], {"name": r["h"] or r["site"] or "(병원명 없음)", "pp": [], "names": []})
        if r["h"] and r["h"] not in h["names"]:
            h["names"].append(r["h"])
        pp = split_people(r["fse"])
        r["pp"] = pp
        for p in pp:
            if p not in h["pp"]:
                h["pp"].append(p)

    matched = len({r["k"] for r in fsn_rows} & fco_keys)

    data = {
        "app": APP_TITLE,
        "version": APP_VERSION,
        "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "today": today.isoformat(),
        "fsn": {
            "available": bool(fsn_parse.get("ok")),
            "error": fsn_parse.get("error", "") if fsn_info else "고객안내문 파일이 없습니다.",
            "file": (fsn_info or {}).get("file", {}),
            "rows": fsn_rows,
            "fields": fsn_parse.get("fields", {}),
            "wide": fsn_parse.get("wide", {}),
            "sheets": fsn_parse.get("sheets", []),
            "statusValues": fsn_parse.get("status_values", []),
            "methodValues": fsn_parse.get("method_values", []),
            "excluded": fsn_parse.get("excluded", 0),
            "skipped": fsn_parse.get("skipped", 0),
            "warnings": fsn_parse.get("warnings", []),
        },
        "fco": {
            "available": any((i.get("parse") or {}).get("ok") for i in fco_infos),
            "sources": [{
                "name": i["source"],
                "ok": bool((i.get("parse") or {}).get("ok")),
                "error": (i.get("parse") or {}).get("error", "") if i.get("parse") else i.get("error", "파일 없음"),
                "file": i.get("file", {}),
                "count": len((i.get("parse") or {}).get("rows", [])),
                "closed": (i.get("parse") or {}).get("closed", 0),
                "sheet": (i.get("parse") or {}).get("sheet", ""),
                "fields": (i.get("parse") or {}).get("fields", {}),
            } for i in fco_infos],
            "rows": fco_rows,
        },
        "hospitals": hospitals,
        "link": {"matchedHospitals": matched, "linkedBySn": linked_by_sn},
    }
    return data


def render_dashboard(data: dict, out_path: Path) -> Path:
    tpl_path = resource_path(TEMPLATE_NAME)
    if not tpl_path.exists():
        raise FileNotFoundError(f"{TEMPLATE_NAME} 파일이 없습니다. main.py와 같은 폴더에 있어야 합니다.")
    tpl = tpl_path.read_text(encoding="utf-8")
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    payload = payload.replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    html = tpl.replace("__DASHBOARD_DATA__", payload, 1)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_suffix(".tmp")
    tmp.write_text(html, encoding="utf-8")
    os.replace(tmp, out_path)
    return out_path


def file_info(path: Path, original_name: str = "", stale: bool = False, origin: str = "") -> dict:
    try:
        mtime = dt.datetime.fromtimestamp(path.stat().st_mtime)
    except OSError:
        mtime = None
    return {
        "name": original_name or path.name,
        "path": str(path),
        "time": mtime.strftime("%Y-%m-%d %H:%M") if mtime else "",
        "stale": stale,
        "origin": origin,
    }


def open_in_browser(path: Path):
    p = str(path)
    if sys.platform == "win32":
        cands = []
        for env in ("PROGRAMFILES(X86)", "PROGRAMFILES", "LOCALAPPDATA"):
            base = os.environ.get(env)
            if base:
                cands.append(Path(base) / "Microsoft" / "Edge" / "Application" / "msedge.exe")
        for c in cands:
            if c.exists():
                try:
                    subprocess.Popen([str(c), p])
                    return
                except Exception:
                    log_exc("Edge 실행 실패")
        try:
            os.startfile(p)  # type: ignore[attr-defined]
            return
        except Exception:
            log_exc("os.startfile 실패")
    webbrowser.open(path.resolve().as_uri())


def open_folder(path: Path):
    try:
        if sys.platform == "win32":
            os.startfile(str(path))  # type: ignore[attr-defined]
        else:
            webbrowser.open(path.resolve().as_uri())
    except Exception:
        log_exc("폴더 열기 실패")


# ---------------------------------------------------------------------------
# 파일 다운로드 (Edge + Selenium)
# ---------------------------------------------------------------------------

def watch_dirs() -> list:
    dirs = [DOWNLOAD_DIR, HOME / "Downloads", HOME / "Desktop"]
    for env in ("OneDrive", "OneDriveCommercial"):
        od = os.environ.get(env)
        if od:
            dirs.append(Path(od) / "Desktop")
    out = []
    for d in dirs:
        if d not in out and d.exists():
            out.append(d)
    return out


def find_new_excel(since: float, name_ok) -> Path | None:
    found = []
    for d in watch_dirs():
        try:
            for f in d.iterdir():
                if not f.is_file() or f.suffix.lower() not in EXCEL_EXTS or f.name.startswith("~$"):
                    continue
                try:
                    st = f.stat()
                except OSError:
                    continue
                if st.st_mtime >= since and st.st_size > 0 and name_ok(f.name):
                    found.append((st.st_mtime, f))
        except OSError:
            continue
    if not found:
        return None
    found.sort(reverse=True)
    f = found[0][1]
    # 다운로드가 끝났는지 (크기 안정 + zip 형식) 확인
    try:
        s1 = f.stat().st_size
        time.sleep(1.0)
        s2 = f.stat().st_size
        if s1 != s2 or not zipfile.is_zipfile(f):
            return None
    except OSError:
        return None
    return f


def wait_for_excel(since: float, timeout: float, name_ok, stop_event=None) -> Path | None:
    end = time.time() + timeout
    while time.time() < end:
        if stop_event is not None and stop_event.is_set():
            return None
        f = find_new_excel(since, name_ok)
        if f:
            return f
        time.sleep(1.5)
    return None


def store_download(src: Path, dest: Path) -> Path:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copyfile(src, dest)
    except PermissionError:
        dest = dest.with_name(f"{dest.stem}_{dt.datetime.now():%H%M%S}{dest.suffix}")
        shutil.copyfile(src, dest)
    os.utime(dest, None)
    meta = {"original_name": src.name, "downloaded_at": dt.datetime.now().isoformat(timespec="seconds")}
    try:
        dest.with_suffix(".meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    except OSError:
        pass
    if src.parent == DOWNLOAD_DIR:
        try:
            src.unlink()
        except OSError:
            pass
    log(f"다운로드 저장: {src} → {dest}")
    return dest


def read_meta(path: Path) -> dict:
    try:
        return json.loads(path.with_suffix(".meta.json").read_text(encoding="utf-8"))
    except Exception:
        return {}


def add_download_param(url: str) -> str:
    if "download=1" in url:
        return url
    return url + ("&" if "?" in url else "?") + "download=1"


FIND_JS = r"""
const cands = arguments[0].map(s => s.trim().toLowerCase());
function vis(el){
  const r = el.getBoundingClientRect();
  if (r.width <= 0 || r.height <= 0) return false;
  const st = getComputedStyle(el);
  return st.visibility !== 'hidden' && st.display !== 'none';
}
const clickSel = 'button,a,[role="menuitem"],[role="menuitemradio"],[role="button"],[role="tab"],[role="option"]';
const els = document.querySelectorAll(clickSel + ',li,span,div,label');
let best = null, bestArea = Infinity;
for (const el of els) {
  const label = (el.getAttribute('aria-label') || el.getAttribute('title') || '').trim().toLowerCase();
  const txt = ((el.innerText || '').trim().toLowerCase().split('\n')[0] || '').trim();
  if (!(cands.includes(txt) || cands.includes(label))) continue;
  if (!vis(el)) continue;
  const r = el.getBoundingClientRect();
  const area = r.width * r.height;
  if (area < bestArea) { best = el; bestArea = area; }
}
if (!best) return null;
return best.closest(clickSel) || best;
"""

LOADING_JS = r"""
const t = ((document.body && document.body.innerText) || '').toLowerCase();
return ['please wait', 'loading the workbook', '잠시 기다려', '통합 문서를 로드', '로드하는 중', '불러오는 중']
  .some(w => t.includes(w));
"""

FETCH_JS = r"""
const done = arguments[arguments.length - 1];
fetch(arguments[0], {credentials: 'include', redirect: 'follow'}).then(async r => {
  const buf = new Uint8Array(await r.arrayBuffer());
  let bin = '';
  for (let i = 0; i < buf.length; i += 0x8000) bin += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000));
  done({ok: r.ok, status: r.status, type: r.headers.get('content-type') || '',
        cd: r.headers.get('content-disposition') || '', url: r.url, size: buf.length, b64: btoa(bin)});
}).catch(e => done({ok: false, error: String(e)}));
"""

FILE_TEXTS = ["File", "파일"]
COPY_TEXTS = ["Create a Copy", "Save As", "Save a Copy", "복사본 만들기", "다른 이름으로 저장", "복사본 저장"]
DOWNLOAD_TEXTS = ["Download a Copy", "복사본 다운로드", "Download a copy to your device", "Download"]


class EdgeDownloader:
    """회사 PC의 Edge(Windows 계정 SSO)로 SharePoint 파일을 받습니다."""

    def __init__(self, report, stop_event):
        self.report = report
        self.stop_event = stop_event
        self.driver = None

    def start(self):
        from selenium import webdriver
        from selenium.webdriver.edge.options import Options
        from selenium.webdriver.edge.service import Service

        PROFILE_DIR.mkdir(parents=True, exist_ok=True)
        DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)
        opts = Options()
        opts.page_load_strategy = "eager"
        opts.add_argument(f"--user-data-dir={PROFILE_DIR}")
        opts.add_argument("--profile-directory=Default")
        opts.add_argument("--no-first-run")
        opts.add_argument("--no-default-browser-check")
        opts.add_argument("--start-maximized")
        opts.add_experimental_option("prefs", {
            "download.default_directory": str(DOWNLOAD_DIR),
            "download.open_pdf_in_system_reader": False,
            "download.extensions_to_open": "",
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "safebrowsing.enabled": True,
            "profile.default_content_setting_values.automatic_downloads": 1,
        })
        opts.add_experimental_option("excludeSwitches", ["enable-logging"])
        service = Service()
        if sys.platform == "win32":
            try:
                service.creation_flags = subprocess.CREATE_NO_WINDOW  # 콘솔 창 숨김
            except Exception:
                pass
        self.report("Edge 브라우저를 여는 중...")
        try:
            self.driver = webdriver.Edge(options=opts, service=service)
        except Exception:
            # 이전 실행이 비정상 종료되어 프로필이 잠긴 경우 등 → 두 번째 프로필로 한 번 더
            log_exc("Edge 시작 실패 → 다른 프로필로 재시도")
            alt = APP_DIR / "browser_profile_2"
            alt.mkdir(parents=True, exist_ok=True)
            opts.arguments[:] = [a for a in opts.arguments if not a.startswith("--user-data-dir=")]
            opts.add_argument(f"--user-data-dir={alt}")
            service2 = Service()
            if sys.platform == "win32":
                try:
                    service2.creation_flags = subprocess.CREATE_NO_WINDOW
                except Exception:
                    pass
            self.driver = webdriver.Edge(options=opts, service=service2)
        self.driver.set_page_load_timeout(45)
        # 다운로드를 묻지 않고 바로 저장 (Edge의 '이 파일을 열까요?' 창 방지)
        for cmd in ("Browser.setDownloadBehavior", "Page.setDownloadBehavior"):
            try:
                self.driver.execute_cdp_cmd(cmd, {"behavior": "allow", "downloadPath": str(DOWNLOAD_DIR)})
                break
            except Exception:
                log(f"{cmd} 실패")
        try:
            self.driver.command_executor.set_timeout(20)
        except Exception:
            try:
                self.driver.command_executor._client_config.timeout = 20  # selenium 신버전
            except Exception:
                pass
        log("Edge 시작")

    def quit(self):
        if self.driver is not None:
            try:
                self.driver.quit()
            except Exception:
                pass
            self.driver = None

    # --- 내부 동작 ---
    def _get(self, url):
        try:
            self.driver.get(url)
        except Exception as e:
            log(f"driver.get 예외(무시하고 계속): {type(e).__name__}: {str(e)[:200]}")

    def _current_url(self):
        try:
            return self.driver.current_url or ""
        except Exception:
            return ""

    def _wait_login(self, label, timeout=150):
        end = time.time() + timeout
        warned = False
        while time.time() < end and not self.stop_event.is_set():
            u = self._current_url().lower()
            if any(w in u for w in ("login.microsoftonline", "login.live", "/adfs/", "login.windows", "sso.")):
                if not warned:
                    self.report(f"{label}: 로그인이 필요합니다. 열린 Edge 창에서 회사 계정으로 로그인해 주세요.")
                    warned = True
                time.sleep(2)
                continue
            return True
        return False

    def _find_in_frames(self, cands, depth=0):
        try:
            el = self.driver.execute_script(FIND_JS, cands)
        except Exception:
            el = None
        if el is not None:
            return el
        if depth >= 3:
            return None
        from selenium.webdriver.common.by import By
        try:
            frames = self.driver.find_elements(By.CSS_SELECTOR, "iframe,frame")
        except Exception:
            frames = []
        for fr in frames:
            try:
                self.driver.switch_to.frame(fr)
            except Exception:
                continue
            el = self._find_in_frames(cands, depth + 1)
            if el is not None:
                return el
            try:
                self.driver.switch_to.parent_frame()
            except Exception:
                break
        return None

    def _find(self, cands, timeout):
        end = time.time() + timeout
        while time.time() < end and not self.stop_event.is_set():
            try:
                self.driver.switch_to.default_content()
            except Exception:
                pass
            el = self._find_in_frames(cands)
            if el is not None:
                return el
            time.sleep(1)
        return None

    def _click(self, el):
        try:
            el.click()
            return True
        except Exception:
            try:
                self.driver.execute_script("arguments[0].scrollIntoView({block:'center'}); arguments[0].click();", el)
                return True
            except Exception:
                log_exc("클릭 실패")
                return False

    def _loading_anywhere(self, depth=0):
        try:
            if self.driver.execute_script(LOADING_JS):
                return True
        except Exception:
            pass
        if depth >= 3:
            return False
        from selenium.webdriver.common.by import By
        try:
            frames = self.driver.find_elements(By.CSS_SELECTOR, "iframe,frame")
        except Exception:
            frames = []
        for fr in frames:
            try:
                self.driver.switch_to.frame(fr)
            except Exception:
                continue
            hit = self._loading_anywhere(depth + 1)
            try:
                self.driver.switch_to.parent_frame()
            except Exception:
                pass
            if hit:
                return True
        return False

    def _wait_workbook_loaded(self, timeout=90):
        end = time.time() + timeout
        clear = 0
        while time.time() < end and not self.stop_event.is_set():
            try:
                self.driver.switch_to.default_content()
            except Exception:
                pass
            if self._loading_anywhere():
                clear = 0
            else:
                clear += 1
                if clear >= 2:
                    return True
            time.sleep(1.5)
        return False

    def _menu_download(self, label) -> bool:
        """Excel Online: 파일 → 복사본 만들기 → 복사본 다운로드"""
        self.report(f"{label}: 엑셀 화면 로딩을 기다리는 중...")
        self._wait_workbook_loaded()
        el = self._find(FILE_TEXTS, 40)
        if el is None:
            log("파일(File) 메뉴를 찾지 못함")
            return False
        self.report(f"{label}: 파일 → 복사본 다운로드 메뉴 클릭 중...")
        if not self._click(el):
            return False
        time.sleep(2)
        el = self._find(COPY_TEXTS + DOWNLOAD_TEXTS[:3], 15)
        if el is None:
            log("복사본 만들기 메뉴를 찾지 못함")
            return False
        txt = ""
        try:
            txt = (el.text or el.get_attribute("aria-label") or "").strip().lower()
        except Exception:
            pass
        self._click(el)
        time.sleep(1.5)
        if not any(t.lower() in txt for t in DOWNLOAD_TEXTS[:3]):
            el = self._find(DOWNLOAD_TEXTS[:3], 15)
            if el is None:
                log("복사본 다운로드 메뉴를 찾지 못함")
                return False
            self._click(el)
        time.sleep(3)
        # 확인 대화상자가 뜨는 경우
        el = self._find(["Download", "다운로드"], 3)
        if el is not None:
            self._click(el)
        return True

    def _fetch_in_page(self, url: str, dest: Path, label: str) -> Path | None:
        """로그인된 Edge 페이지 안에서 fetch로 파일 내용을 직접 가져옴.
        브라우저 다운로드를 쓰지 않으므로 '이 파일을 열까요?' 창이 뜨지 않습니다."""
        from urllib.parse import urlparse, unquote
        self.report(f"{label}: 로그인 확인 중...")
        self._get(url)
        self._wait_login(label)
        host = urlparse(url).netloc.lower()
        end = time.time() + 60
        while time.time() < end and not self.stop_event.is_set():
            if urlparse(self._current_url()).netloc.lower() == host:
                break
            time.sleep(1)
        self.report(f"{label}: 파일 받는 중...")
        try:
            self.driver.switch_to.default_content()
        except Exception:
            pass
        try:
            self.driver.set_script_timeout(240)
            res = self.driver.execute_async_script(FETCH_JS, add_download_param(url))
        except Exception as e:
            log(f"{label}: fetch 실패 {type(e).__name__}: {str(e)[:200]}")
            return None
        if not isinstance(res, dict) or not res.get("b64"):
            log(f"{label}: fetch 결과 없음 {str(res)[:300] if not isinstance(res, dict) else {k: v for k, v in res.items() if k != 'b64'}}")
            return None
        import base64
        data = base64.b64decode(res["b64"])
        info = {k: v for k, v in res.items() if k != "b64"}
        if not data.startswith(b"PK"):
            log(f"{label}: fetch 결과가 엑셀이 아님 {info}")
            return None
        name = ""
        cd = res.get("cd") or ""
        m = re.search(r"filename\*=(?:UTF-8|utf-8)''([^;]+)", cd) or re.search(r'filename="?([^";]+)"?', cd)
        if m:
            name = unquote(m.group(1)).strip()
        name = re.sub(r'[\\/:*?"<>|]+', "_", name) or "download.xlsx"
        if not name.lower().endswith(EXCEL_EXTS):
            name += ".xlsx"
        DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)
        tmp = DOWNLOAD_DIR / name
        tmp.write_bytes(data)
        log(f"{label}: fetch 성공 {len(data):,} bytes, {name}")
        return store_download(tmp, dest)

    def download(self, url: str, dest: Path, name_ok, label: str, allow_manual=True) -> Path | None:
        if self.driver is None:
            self.start()
        since = time.time() - 2
        # 0) 페이지 안에서 직접 받기 (다운로드 창 없음)
        got = self._fetch_in_page(url, dest, label)
        if got:
            return got
        # 1) 공유 링크 + download=1 → 브라우저 다운로드
        self.report(f"{label}: 파일 받는 중 (바로 받기)...")
        self._get(add_download_param(url))
        self._wait_login(label)
        f = wait_for_excel(since, 35, name_ok, self.stop_event)
        if f:
            return store_download(f, dest)
        # 2) Excel Online 메뉴 클릭 (최대 3회)
        for attempt in range(1, 4):
            if self.stop_event.is_set():
                return None
            log(f"{label}: 메뉴 다운로드 시도 {attempt}")
            self._get(url)
            self._wait_login(label)
            ok = self._menu_download(label)
            if ok:
                self.report(f"{label}: 다운로드 완료를 기다리는 중...")
                f = wait_for_excel(since, 60, name_ok, self.stop_event)
                if f:
                    return store_download(f, dest)
        # 3) 수동
        if allow_manual:
            self.report(f"{label}: 자동 다운로드가 되지 않았습니다.\n"
                        "열린 Edge 창에서 [파일 → 복사본 만들기 → 복사본 다운로드]를 직접 눌러 주세요. (3분 동안 기다립니다)")
            f = wait_for_excel(since, 180, name_ok, self.stop_event)
            if f:
                return store_download(f, dest)
        return None


# ---------------------------------------------------------------------------
# 전체 처리 흐름
# ---------------------------------------------------------------------------

def is_today(path: Path) -> bool:
    try:
        return dt.date.fromtimestamp(path.stat().st_mtime) == dt.date.today()
    except OSError:
        return False


def fco_name_ok(source: str):
    pat = re.compile(FCO_SOURCES[source]["pattern"])

    def ok(name: str) -> bool:
        return "fco" in name.lower() and bool(pat.search(name))
    return ok


def fsn_name_ok(name: str) -> bool:
    # 고객안내문 파일 이름은 모르므로, 받기 시작한 뒤 새로 생긴 엑셀이면 모두 인정
    return True


def run_pipeline(force: bool, manual_fsn: Path | None, report, step, stop_event, holder: dict) -> Path:
    """report(text), step(idx, state, text). 대시보드 HTML 경로 반환."""
    ensure_dirs()
    cfg = load_config()
    today = dt.date.today()
    fsn_url = cfg.get("fsn_url") or FSN_SOURCE["url"]
    fco_urls = {k: (cfg.get("fco_urls") or {}).get(k) or v["url"] for k, v in FCO_SOURCES.items()}
    col_overrides = cfg.get("columns") or None

    # ---------- 1. 고객안내문 파일 ----------
    fsn_cache = DATA_DIR / FSN_SOURCE["cache"]
    fsn_path, fsn_stale, fsn_origin = None, False, ""
    need_fsn_download = False
    if manual_fsn:
        step(0, "run", "선택한 파일 복사 중")
        fsn_path = store_download(Path(manual_fsn), fsn_cache) if Path(manual_fsn) != fsn_cache else fsn_cache
        fsn_origin = "직접 선택"
        step(0, "ok", f"직접 선택: {Path(manual_fsn).name}")
    elif fsn_cache.exists() and is_today(fsn_cache) and not force:
        fsn_path, fsn_origin = fsn_cache, "오늘 받은 파일"
        step(0, "ok", "오늘 받은 파일 사용")
    else:
        need_fsn_download = True

    # ---------- 2. FCO 파일 ----------
    fco_plan = []  # (source, path or None, stale, origin, need_download)
    for src, spec in FCO_SOURCES.items():
        mine = DATA_DIR / spec["cache"]
        theirs = FCO_APP_DIR / spec["cache"]
        today_files = [p for p in (mine, theirs) if p.exists() and is_today(p)]
        if today_files and not force:
            fco_plan.append([src, today_files[0], False,
                             "오늘 받은 파일" if today_files[0] == mine else "FCO 프로그램이 오늘 받은 파일", False])
        elif fco_urls.get(src):
            fco_plan.append([src, None, False, "", True])
        else:
            existing = sorted([p for p in (mine, theirs) if p.exists()], key=lambda p: p.stat().st_mtime,
                              reverse=True)
            if existing:
                fco_plan.append([src, existing[0], True, "이전에 받은 파일", False])
            else:
                fco_plan.append([src, None, False, "", False])

    downloads_needed = need_fsn_download or any(p[4] for p in fco_plan)
    downloader = None
    fsn_fail_reason = ""
    try:
        if downloads_needed:
            downloader = EdgeDownloader(report, stop_event)
            holder["downloader"] = downloader
            try:
                downloader.start()
            except Exception as e:
                log_exc("Edge 시작 실패")
                fsn_fail_reason = f"Edge 브라우저를 열지 못했습니다 ({type(e).__name__})"
                report(fsn_fail_reason)
                downloader = None

        if need_fsn_download:
            step(0, "run", "SharePoint에서 받는 중")
            got = None
            if downloader is not None:
                try:
                    got = downloader.download(fsn_url, fsn_cache, fsn_name_ok, "고객안내문")
                    if not got:
                        fsn_fail_reason = "다운로드된 파일을 찾지 못했습니다 (시간 초과)"
                except Exception as e:
                    log_exc("고객안내문 다운로드 실패")
                    fsn_fail_reason = f"다운로드 중 오류 ({type(e).__name__})"
            if got:
                fsn_path, fsn_origin = got, "방금 받은 파일"
                step(0, "ok", "새로 받음")
            elif fsn_cache.exists():
                fsn_path, fsn_stale, fsn_origin = fsn_cache, True, "이전에 받은 파일"
                d = dt.datetime.fromtimestamp(fsn_cache.stat().st_mtime).strftime("%m/%d %H:%M")
                step(0, "warn", f"받기 실패 → 이전 파일 사용 ({d})")
            else:
                step(0, "fail", "받기 실패")

        step(1, "run", "확인 중")
        for item in fco_plan:
            src, path, stale, origin, need = item
            if need:
                spec = FCO_SOURCES[src]
                got = None
                if downloader is not None:
                    try:
                        got = downloader.download(fco_urls[src], DATA_DIR / spec["cache"], fco_name_ok(src),
                                                  f"FCO {src}", allow_manual=False)
                    except Exception:
                        log_exc(f"FCO {src} 다운로드 실패")
                if got:
                    item[1], item[3] = got, "방금 받은 파일"
                else:
                    existing = sorted([p for p in (DATA_DIR / spec["cache"], FCO_APP_DIR / spec["cache"])
                                       if p.exists()], key=lambda p: p.stat().st_mtime, reverse=True)
                    if existing:
                        item[1], item[2], item[3] = existing[0], True, "받기 실패 → 이전 파일"
    finally:
        if downloader is not None:
            downloader.quit()
        holder["downloader"] = None

    fco_ok = [p for p in fco_plan if p[1]]
    fco_missing = [p[0] for p in fco_plan if not p[1]]
    if not fco_ok:
        step(1, "warn", "FCO 파일 없음 (FCO 없이 진행)")
    elif any(p[2] for p in fco_ok) or fco_missing:
        def _d(path):
            return dt.datetime.fromtimestamp(path.stat().st_mtime).strftime("%m/%d")
        txt = " · ".join(f"{p[0]} {_d(p[1]) if p[2] else '오늘'}" for p in fco_ok)
        if fco_missing:
            txt += f"  ({', '.join(fco_missing)} 없음)"
        if not any(fco_urls.values()):
            txt += "\n→ FCO 주소 미설정: FCO 프로그램이 받아 둔 파일 사용"
        step(1, "warn", txt)
    else:
        step(1, "ok", " · ".join(p[0] for p in fco_ok))

    # ---------- 3. 대시보드 ----------
    if stop_event.is_set():
        raise RuntimeError("취소되었습니다.")
    step(2, "run", "데이터 정리 중")
    report("엑셀 데이터를 읽는 중...")
    fsn_info = None
    if fsn_path:
        meta = read_meta(fsn_path)
        report("엑셀 데이터를 읽는 중... (파일이 크면 1~2분 걸릴 수 있습니다)")
        fsn_info = {"parse": parse_fsn(fsn_path, col_overrides, cfg.get("skip_sheets")),
                    "file": file_info(fsn_path, meta.get("original_name", ""), fsn_stale, fsn_origin)}
    fco_infos = []
    for src, path, stale, origin, _ in fco_plan:
        if path:
            meta = read_meta(path)
            fco_infos.append({"source": src, "parse": parse_fco(path, src, today),
                              "file": file_info(path, meta.get("original_name", ""), stale, origin)})
        else:
            fco_infos.append({"source": src, "parse": None, "file": {},
                              "error": "파일 없음 (FCO 주소 미설정, FCO 프로그램 캐시도 없음)"})

    if not (fsn_info and fsn_info["parse"].get("ok")) and not any((i.get("parse") or {}).get("ok")
                                                                  for i in fco_infos):
        if fsn_info and fsn_info["parse"].get("error"):
            msg = "고객안내문 파일을 읽지 못했습니다.\n" + fsn_info["parse"]["error"]
        else:
            msg = ("고객안내문 파일을 받지 못했습니다.\n원인: " + (fsn_fail_reason or "알 수 없음") +
                   "\n\n해결 방법\n1) [새로 받기]를 눌러 다시 시도\n"
                   "2) 브라우저에서 파일을 직접 내려받은 뒤 [파일 직접 선택]으로 선택")
        raise RuntimeError(msg)

    data = build_dashboard_data(fsn_info, fco_infos, today)
    if not fsn_info:
        data["fsn"]["error"] = "고객안내문 파일을 받지 못했습니다. " + (fsn_fail_reason or "")
    out = render_dashboard(data, DASHBOARD_PATH)
    n_fsn = len(data["fsn"]["rows"])
    n_fco = len(data["fco"]["rows"])
    if fsn_info and not fsn_info["parse"].get("ok"):
        step(2, "warn", f"고객안내문 읽기 실패 · FCO {n_fco}건")
    else:
        step(2, "ok", f"고객안내문 {n_fsn}건 · FCO {n_fco}건")
    log(f"대시보드 생성: {out}")
    return out


# ---------------------------------------------------------------------------
# GUI (Tkinter)
# ---------------------------------------------------------------------------

def run_gui():
    import tkinter as tk
    from tkinter import ttk, messagebox, filedialog

    if sys.platform == "win32":
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass

    root = tk.Tk()
    root.title(APP_TITLE)
    try:
        ico = resource_path("FSN.ico")
        if ico.exists():
            root.iconbitmap(str(ico))
    except Exception:
        pass

    font_family = "Malgun Gothic" if sys.platform == "win32" else "TkDefaultFont"
    f_title = (font_family, 14, "bold")
    f_body = (font_family, 10)
    f_small = (font_family, 9)

    q: queue.Queue = queue.Queue()
    stop_event = threading.Event()
    holder = {"downloader": None, "running": False, "html": None}

    outer = tk.Frame(root, padx=22, pady=18, bg="#ffffff")
    outer.pack(fill="both", expand=True)
    root.configure(bg="#ffffff")

    btns = tk.Frame(outer, bg="#ffffff")
    btns.pack(side="bottom", fill="x", pady=(14, 0))

    tk.Label(outer, text="고객안내문 · FCO 통합 Dashboard", font=f_title, bg="#ffffff", fg="#0b0b0b",
             anchor="w").pack(fill="x")
    tk.Label(outer, text="SharePoint 파일을 받아서 브라우저 대시보드로 보여줍니다.", font=f_small, bg="#ffffff",
             fg="#52514e", anchor="w").pack(fill="x", pady=(2, 12))

    steps_frame = tk.Frame(outer, bg="#ffffff")
    steps_frame.pack(fill="x")
    step_names = ["고객안내문 파일", "FCO 파일 (MR · XR · CT/AMI)", "대시보드 만들기"]
    step_icons, step_texts, step_state = [], [], [None, None, None]
    for i, name in enumerate(step_names):
        row = tk.Frame(steps_frame, bg="#ffffff")
        row.pack(fill="x", pady=3)
        ic = tk.Label(row, text="○", width=2, font=(font_family, 11, "bold"), bg="#ffffff", fg="#898781")
        ic.pack(side="left")
        tk.Label(row, text=f"{i + 1}. {name}", font=f_body, bg="#ffffff", fg="#0b0b0b", width=26,
                 anchor="w").pack(side="left")
        st = tk.Label(row, text="대기", font=f_small, bg="#ffffff", fg="#52514e", anchor="w", justify="left")
        st.pack(side="left", fill="x", expand=True)
        step_icons.append(ic)
        step_texts.append(st)

    pb = ttk.Progressbar(outer, mode="indeterminate")
    pb.pack(fill="x", pady=(14, 6))
    status_var = tk.StringVar(value="준비 중...")
    tk.Label(outer, textvariable=status_var, font=f_small, bg="#ffffff", fg="#52514e", justify="left",
             anchor="w", wraplength=500).pack(fill="x")

    def set_buttons(running: bool):
        for b in (b_open, b_refresh, b_pick):
            b.configure(state="disabled" if running else "normal")
        if not holder.get("html"):
            b_open.configure(state="disabled")

    def start(force=False, manual=None):
        if holder["running"]:
            return
        holder["running"] = True
        stop_event.clear()
        for i in range(3):
            step_state[i] = None
            step_icons[i].configure(text="○", fg="#898781")
            step_texts[i].configure(text="대기", fg="#52514e")
        set_buttons(True)
        pb.configure(mode="indeterminate", value=0)
        pb.start(12)

        def worker():
            try:
                html = run_pipeline(force, manual, lambda t: q.put(("status", t)),
                                    lambda i, s, t: q.put(("step", i, s, t)), stop_event, holder)
                q.put(("done", html))
            except Exception as e:
                log_exc("처리 실패")
                q.put(("error", str(e)))

        threading.Thread(target=worker, daemon=True).start()

    def on_open():
        if holder.get("html"):
            open_in_browser(holder["html"])

    def on_pick():
        p = filedialog.askopenfilename(title="고객안내문 엑셀 파일 선택",
                                       filetypes=[("Excel", "*.xlsx *.xlsm"), ("All", "*.*")])
        if p:
            start(force=False, manual=Path(p))

    def on_close():
        stop_event.set()
        d = holder.get("downloader")
        if d is not None:
            threading.Thread(target=d.quit, daemon=True).start()
        root.after(300, root.destroy)

    b_open = ttk.Button(btns, text="대시보드 열기", command=on_open)
    b_refresh = ttk.Button(btns, text="새로 받기", command=lambda: start(force=True))
    b_pick = ttk.Button(btns, text="파일 직접 선택", command=on_pick)
    b_folder = ttk.Button(btns, text="폴더 열기", command=lambda: open_folder(APP_DIR))
    for b in (b_open, b_refresh, b_pick, b_folder):
        b.pack(side="left", padx=(0, 6))
    ttk.Button(btns, text="닫기", command=on_close).pack(side="right")

    STATE_STYLE = {"run": ("◐", "#2a78d6"), "ok": ("✔", "#006300"), "warn": ("!", "#b26a00"),
                   "fail": ("✖", "#d03b3b")}

    def poll():
        try:
            while True:
                msg = q.get_nowait()
                kind = msg[0]
                if kind == "status":
                    status_var.set(msg[1])
                elif kind == "step":
                    _, i, s, t = msg
                    step_state[i] = s
                    sym, color = STATE_STYLE.get(s, ("○", "#898781"))
                    step_icons[i].configure(text=sym, fg=color)
                    step_texts[i].configure(text=t, fg="#0b0b0b" if s != "fail" else "#d03b3b")
                elif kind == "done":
                    holder["running"] = False
                    holder["html"] = msg[1]
                    pb.stop()
                    pb.configure(mode="determinate", value=100)
                    status_var.set(f"완료 · 대시보드를 브라우저로 열었습니다.\n{msg[1]}")
                    set_buttons(False)
                    open_in_browser(msg[1])
                elif kind == "error":
                    holder["running"] = False
                    pb.stop()
                    pb.configure(mode="determinate", value=0)
                    for i, s in enumerate(step_state):
                        if s == "run":
                            step_icons[i].configure(text="✖", fg="#d03b3b")
                            step_texts[i].configure(text="중단됨", fg="#d03b3b")
                    status_var.set("완료하지 못했습니다. (자세한 내용: debug_log.txt)")
                    set_buttons(False)
                    messagebox.showerror(APP_TITLE, f"{msg[1]}\n\n로그: {LOG_PATH}")
        except queue.Empty:
            pass
        root.after(150, poll)

    def report_callback_exception(exc, val, tb):
        log("Tk 콜백 오류\n" + "".join(traceback.format_exception(exc, val, tb)))
        messagebox.showerror(APP_TITLE, f"오류: {val}\n\n로그: {LOG_PATH}")

    root.report_callback_exception = report_callback_exception
    root.protocol("WM_DELETE_WINDOW", on_close)
    root.update_idletasks()
    w, h = max(600, root.winfo_reqwidth()), root.winfo_reqheight()
    x = (root.winfo_screenwidth() - w) // 2
    y = (root.winfo_screenheight() - h) // 3
    root.geometry(f"+{x}+{y}")  # 크기는 내용에 맞춰 자동으로 늘어나게 둠
    root.minsize(w, h)
    root.after(150, poll)
    root.after(300, lambda: start(force=False))
    root.mainloop()


# ---------------------------------------------------------------------------
# 진입점
# ---------------------------------------------------------------------------

def run_cli(args):
    today = dt.date.fromisoformat(args.today) if args.today else dt.date.today()
    fsn_info = None
    if args.build:
        p = Path(args.build)
        fsn_info = {"parse": parse_fsn(p), "file": file_info(p, origin="직접 지정")}
    fco_infos = []
    for spec in args.fco or []:
        src, _, path = spec.partition("=")
        p = Path(path)
        fco_infos.append({"source": src, "parse": parse_fco(p, src, today), "file": file_info(p, origin="직접 지정")})
    data = build_dashboard_data(fsn_info, fco_infos, today)
    if args.sample:
        data["sample"] = True
    out = render_dashboard(data, Path(args.out))
    print(f"FSN rows={len(data['fsn']['rows'])} FCO rows={len(data['fco']['rows'])} → {out}")
    if args.json:
        Path(args.json).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    if not args.no_open:
        open_in_browser(out)


def main():
    # --noconsole 빌드에서는 stdout/stderr가 None이라 print가 오류를 낼 수 있음
    if sys.stdout is None or sys.stderr is None:
        devnull = open(os.devnull, "w", encoding="utf-8")
        sys.stdout = sys.stdout or devnull
        sys.stderr = sys.stderr or devnull

    ap = argparse.ArgumentParser(description=APP_TITLE)
    ap.add_argument("--build", help="(개발용) 고객안내문 엑셀 경로 → HTML 생성")
    ap.add_argument("--fco", action="append", help="(개발용) SOURCE=경로, 예: MR=MR.xlsx")
    ap.add_argument("--out", default=str(DASHBOARD_PATH))
    ap.add_argument("--json", help="(개발용) 데이터 JSON 저장 경로")
    ap.add_argument("--today", help="(개발용) 기준일 YYYY-MM-DD")
    ap.add_argument("--sample", action="store_true", help="(개발용) 샘플 데이터 표시")
    ap.add_argument("--no-open", action="store_true")
    args = ap.parse_args()

    if args.build or args.fco:
        run_cli(args)
        return
    try:
        log(f"===== 시작 v{APP_VERSION} (Python {sys.version.split()[0]}) =====")
        run_gui()
    except Exception as e:
        log_exc("치명적 오류")
        try:
            import tkinter as tk
            from tkinter import messagebox
            r = tk.Tk()
            r.withdraw()
            messagebox.showerror(APP_TITLE, f"프로그램 오류: {e}\n\n로그: {LOG_PATH}")
            r.destroy()
        except Exception:
            pass


if __name__ == "__main__":
    main()
