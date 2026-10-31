# -*- coding: utf-8 -*-
"""
간단 회귀 테스트 (pytest 없이 실행 가능):  python tests/test_parsing.py
샘플 엑셀을 임시 폴더에 만들어 파싱/대시보드 생성까지 확인합니다.
"""
import datetime as dt
import json
import os
import sys
import tempfile
import threading
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))

import main  # noqa: E402
import make_sample_data  # noqa: E402


def test_classify_status():
    cases = {
        "완료": "done", "전달완료": "done", "O": "done", "o": "done", "Y": "done", "발송": "done", "등기": "done",
        "이메일 발송": "done", "회신완료": "done", "2026-10-01": "done", dt.date(2026, 10, 1): "done",
        "미전달": "todo", "미발송": "todo", "X": "todo", "N": "todo", "진행중": "todo", "예정": "todo",
        "대기": "todo", "반송": "todo", "-": "",  # "-"는 빈칸 취급 → 미전달
        "해당없음": "excluded", "N/A": "excluded", "철거": "excluded",
        "": "", None: "",
        "???": "unknown",
    }
    for v, want in cases.items():
        got = main.classify_status(v)
        assert got == want, f"classify_status({v!r}) = {got!r}, 기대값 {want!r}"


def test_canon_method():
    cases = {
        "등기": "등기", "등기우편": "등기", "우편": "등기", "직접전달": "직접전달", "방문": "직접전달",
        "이메일": "이메일", "E-mail": "이메일", "메일": "이메일", "등기, 이메일": "등기+이메일",
        "Registered mail": "등기", "": "", "미정": "",
    }
    for v, want in cases.items():
        got = main.canon_method(v)
        assert got == want, f"canon_method({v!r}) = {got!r}, 기대값 {want!r}"


def test_to_date():
    assert main.to_date("2026-10-23") == dt.date(2026, 10, 23)
    assert main.to_date("2026.10.3") == dt.date(2026, 10, 3)
    assert main.to_date("2026년 10월 3일") == dt.date(2026, 10, 3)
    assert main.to_date("26.10.03") == dt.date(2026, 10, 3)
    assert main.to_date(46318) == dt.date(2026, 10, 23)  # 엑셀 일련번호
    assert main.to_date("완료") is None
    assert main.to_date(None) is None


def test_hospital_key():
    assert main.hospital_key("(재)가람 의료원") == main.hospital_key("가람의료원")
    assert main.hospital_key("의료법인 누리병원") == main.hospital_key("누리 병원")
    assert main.hospital_key("한빛대학교병원") != main.hospital_key("한빛병원")


def test_match_columns():
    hdr = [(0, "No"), (1, "병원명"), (2, "전달 방법"), (3, "전달여부"), (4, "전달일"), (5, "C&R"), (6, "BU"),
           (7, "담당 FSE"), (8, "영업담당"), (9, "S/N"), (10, "고객안내문 회신여부")]
    m = main.match_columns(hdr, main.FSN_FIELDS)
    want = {"hospital": 1, "method": 2, "status": 3, "date": 4, "cr": 5, "bu": 6, "fse": 7, "sales": 8, "sn": 9,
            "reply": 10}
    for k, v in want.items():
        assert m.get(k) == v, f"{k}: {m.get(k)} != {v}  (전체 {m})"


def _samples(tmp: Path):
    make_sample_data.sample_a(tmp / "A.xlsx")
    make_sample_data.sample_b(tmp / "B.xlsx")
    make_sample_data.fco_file(tmp / "MR FCO Action Plan.xlsx", "MR", ["MR"])


def test_parse_samples():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        _samples(tmp)
        a = main.parse_fsn(tmp / "A.xlsx")
        assert a["ok"], a["error"]
        assert a["fields"]["hospital"] == "병원명" and a["fields"]["status"] == "전달여부"
        assert [s["name"] for s in a["sheets"] if s["used"]] == ["전달현황"], "요약 시트는 제외되어야 함"
        assert a["excluded"] > 0, "해당없음 행은 집계 제외"
        assert all(r["h"] for r in a["rows"])

        b = main.parse_fsn(tmp / "B.xlsx")
        assert b["ok"], b["error"]
        assert set(b["wide"].values()) == {"등기", "직접전달", "이메일"}, b["wide"]
        assert all(r["h"] for r in b["rows"]), "세로 병합된 병원명이 채워져야 함"
        assert all((r["s"] == 1) == bool(r["m"]) for r in b["rows"]), "전달방법 표시가 있으면 완료"

        today = dt.date(2026, 10, 23)
        f = main.parse_fco(tmp / "MR FCO Action Plan.xlsx", "MR", today)
        assert f["ok"], f["error"]
        assert f["closed"] > 0 and all(r["st"].lower() != "closed" for r in f["rows"])
        assert {r["b"] for r in f["rows"]} <= {"r", "o", "g", "n"}

        data = main.build_dashboard_data({"parse": a, "file": {}},
                                         [{"source": "MR", "parse": f, "file": {}}], today)
        assert data["link"]["matchedHospitals"] > 0
        out = main.render_dashboard(data, tmp / "out.html")
        html = out.read_text(encoding="utf-8")
        assert "__DASHBOARD_DATA__" not in html and "</script>" in html
        start = html.index('id="dashboard-data" type="application/json">') + len('id="dashboard-data" type="application/json">')
        payload = html[start:html.index("</script>", start)]
        assert json.loads(payload)["fsn"]["available"] is True


def test_sheet_per_fsn():
    """FSN별로 시트가 나뉜 경우: 시트 이름을 FSN으로, 요약 시트는 제외"""
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "multi.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "FSN-001"
        ws.append(["병원명", "BU", "전달여부", "비고"])
        ws.append(["가람병원", "MR", "완료", ""])
        ws.append(["누리병원", "MR", "", "재방문"])
        ws2 = wb.create_sheet("FSN-002")
        ws2.append(["병원명", "BU", "전달여부"])
        ws2.append(["가람병원", "CT", "미전달"])
        ws3 = wb.create_sheet("요약")
        ws3.append(["병원명", "전달여부"])
        ws3.append(["가람병원", "1/2"])
        wb.save(path)
        r = main.parse_fsn(path)
        assert r["ok"], r["error"]
        used = [s["name"] for s in r["sheets"] if s["used"]]
        assert used == ["FSN-001", "FSN-002"], used
        assert sorted({x["fsn"] for x in r["rows"]}) == ["FSN-001", "FSN-002"]
        assert len(r["rows"]) == 3 and sum(x["s"] for x in r["rows"]) == 1


def test_real_layout_korean_names():
    """실제 파일 화면(2026-10-23 캡처)과 비슷한 구조:
    시트 MR/CT/IGT…(BU 컬럼 없음), 영문병원명 + 국문병원명 + 고객명(국문), (삭제예정)/Closed 시트"""
    import openpyxl
    head = ["C&R#", "장비 ID", "장비코드(6NC/12NC)", "장비명", "일련번호(S/N)", "고객코드(향후 업데이트)",
            "영문병원명(업데이트예정)", "국문병원명", "고객명(국문)", "전달방법", "최초 발송일", "지역", "비고",
            "최종증빙 확인", "수신확인서 (Y/N/NA)"]
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "real.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "IB_raw data"
        ws.append(["장비 ID", "영문병원명", "국문병원명", "장비명"])
        ws.append(["1", "A HOSP", "가람병원", "Ingenia"])
        for title in ("MR", "CT", "(삭제예정)SMAX자료", "Closed"):
            w = wb.create_sheet(title)
            w.append(head)
            w.append(["2026-PD-MR-004", "98162875", "782107", "Ingenia", "62030", "0", "DONGKOOK UNIV",
                      "동국대학교병원", "", "등기", dt.date(2026, 9, 1), "Central", "", "O", "Y"])
            w.append(["2026-PD-MR-004", "98747245", "782107", "Ingenia", "62039", "0", "-",
                      "-", "혜민병원", "직접전달", None, "Central", "", "증빙 스캔본", "N"])
            w.append(["2026-PD-MR-018", "100342128", "782107", "Ingenia", "62073", "0", "KOREA UNIV",
                      "", "", "", dt.date(2026, 9, 3), "JR", "", "", "NA"])
        wb.save(path)
        r = main.parse_fsn(path)
        assert r["ok"], r["error"]
        assert r["fields"]["hospital"] == "국문병원명", r["fields"]
        assert r["fields"]["hospital2"] == "고객명(국문)", r["fields"]
        assert r["fields"]["hospital_en"] == "영문병원명(업데이트예정)", r["fields"]
        assert r["fields"]["sn"] == "일련번호(S/N)" and r["fields"]["cr"] == "C&R#", r["fields"]
        used = [s["name"] for s in r["sheets"] if s["used"]]
        assert used == ["MR", "CT"], used
        names = [x["h"] for x in r["rows"] if x["bu"] == "MR"]
        assert names == ["동국대학교병원", "혜민병원", "KOREA UNIV"], names
        assert {x["bu"] for x in r["rows"]} == {"MR", "CT"}
        # 완료 = 최종증빙 확인 칸에 값이 있음 (발송일만 있고 증빙이 비면 미전달)
        assert r["fields"]["status"] == "최종증빙 확인", r["fields"]
        assert [x["s"] for x in r["rows"] if x["bu"] == "MR"] == [1, 1, 0]


def _make_pm(path):
    """PM Planning 파일 모양 가정: Hospital table(병원 × 모달리티별 FSE) + FSE 명단(영문명 → 한글 이름)"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PM Plan 2025"
    ws.append(["S/N", "Jan", "Feb"])
    ws.append(["62030", "PM", ""])
    h = wb.create_sheet("Hospital table")
    h.append(["2025 담당 병원표"])
    h.append(["고객코드", "병원명", "영문병원명", "지역", "MR", "CT", "XR", "Sales"])
    h.append(["1001", "동국대학교병원", "DONGKOOK UNIV", "Central", "김민준", "이서연", "", "윤서준"])
    h.append(["1002", "혜민병원", "HYEMIN HOSP", "Central", "Minjun Kim", "", "정하늘", "장미래"])
    h.append(["1003", "", "KOREA UNIV", "JR", "Jihoon Park", "최유진", "", ""])
    f = wb.create_sheet("FSE")
    f.append(["이름", "영문명", "사번", "Modality", "지역"])
    f.append(["김민준", "Minjun Kim", "K001", "MR", "Central"])
    f.append(["박지훈", "Jihoon Park", "K002", "MR", "JR"])
    f.append(["이서연", "Seoyeon Lee", "K003", "CT", "Central"])
    wb.save(path)


def test_people_mapping():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        _make_pm(tmp / "pm.xlsx")
        pp = main.parse_people(tmp / "pm.xlsx")
        assert pp["ok"], pp["error"]
        used = [x["name"] for x in pp["sheets"] if x["used"]]
        assert used == ["Hospital table"], pp["sheets"]
        assert pp["aliases"] >= 3
        rows = [
            {"h": "동국대학교병원", "hen": "", "bu": "MR", "sn": "", "eq": "", "code": "", "fse": "", "sales": ""},
            {"h": "동국대학교병원", "hen": "", "bu": "CT", "sn": "", "eq": "", "code": "", "fse": "", "sales": ""},
            {"h": "혜민", "hen": "", "bu": "MR", "sn": "", "eq": "", "code": "1002", "fse": "", "sales": ""},
            {"h": "KOREA UNIV", "hen": "", "bu": "MR", "sn": "", "eq": "", "code": "", "fse": "", "sales": ""},
            {"h": "없는병원", "hen": "", "bu": "MR", "sn": "", "eq": "", "code": "", "fse": "", "sales": ""},
        ]
        res = main.apply_people(rows, [], None, pp)
        assert rows[0]["fse"] == "김민준" and rows[0]["sales"] == "윤서준", rows[0]
        assert rows[1]["fse"] == "이서연", rows[1]
        assert rows[2]["fse"] == "김민준", rows[2]          # 고객코드로 연결 + 영문명 → 한글 이름
        assert rows[3]["fse"] == "박지훈", rows[3]          # 영문병원명으로 연결
        assert rows[4]["fse"] == "" and res["stats"]["none"] == 1
        assert res["stats"]["code"] == 1 and res["stats"]["name"] == 3


def test_people_matching_robust():
    """실제 파일에서 담당자가 거의 비던 문제(2026-10-30) 대응 확인"""
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "pm2.xlsx"
        wb = openpyxl.Workbook()
        f = wb.active
        f.title = "FSE"
        f.append(["FSE", "담당 병원"])
        f.append(["김광수", "서울대학교병원, 가람병원\n누리병원"])     # 한 칸에 병원 여러 개
        f.append(["오요한", "삼성서울병원(강남)"])
        f.append(["담당없음", "도담병원"])
        f.append(["박기표", "한빛병원"])
        for nm in ("신동민", "서태훈", "조찬민", "성경주", "설재민", "임형묵"):  # FSE 명단 (병원 칸 비어 있음)
            f.append([nm, ""])
        p = wb.create_sheet("CT PM Plan")                               # 이름은 대상 아니지만 S/N 있는 장비 시트
        p.append([None, None, "FSE", None, "Customer info", None, "기타"])
        p.append(["S/N", "병원", "정", "부", "Name", "연락처", "확인자"])
        p.append(["CT100100", "아무병원", "신동민", "서태훈", "홍고객", "010", "가길동"])
        p.append(["CT100101", "아무병원", "조찬민", "성경주", "김고객", "010", "나길동"])
        p.append(["CT100102", "아무병원", "설재민", "", "이고객", "010", "다길동"])
        p.append(["CT100103", "아무병원", "임형묵", "", "박고객", "010", "라길동"])
        wb.save(path)
        pp = main.parse_people(path)
        assert pp["ok"], pp
        used = sorted(x["name"] for x in pp["sheets"] if x["used"])
        assert used == ["CT PM Plan", "FSE"], pp["sheets"]
        ct = next(x for x in pp["sheets"] if x["name"] == "CT PM Plan")
        assert sorted(x["header"] for x in ct["people"]) == ["부", "정"], ct      # 고객 쪽 이름 칸은 제외
        assert sorted(x["header"] for x in ct["dropped"]) == ["Name", "확인자"], ct
        assert pp["roster"] == 9 and len(pp["names"]) <= 9, (pp["roster"], pp["names"])
        def row(h, bu="CT", sn=""):
            return {"h": h, "hen": "", "bu": bu, "sn": sn, "eq": "", "code": "", "fse": "", "sales": ""}
        rows = [row("서울대병원"), row("누리병원"), row("삼성서울병원"), row("도담병원"), row("모름", sn="00CT100100"),
                row("하람병원", bu="MR"), row("하람병원", bu="CT")]
        res = main.apply_people(rows, [], None, pp)
        assert rows[0]["fse"] == "김광수", rows[0]                # 서울대학교병원 ~ 서울대병원
        assert rows[1]["fse"] == "김광수", rows[1]                # 여러 병원이 적힌 칸
        assert rows[2]["fse"] == "오요한", rows[2]                # 괄호 지점명 무시
        assert rows[3]["fse"] == "", rows[3]                      # "담당없음"은 사람 아님
        assert rows[4]["fse"] == "신동민, 서태훈", rows[4]          # S/N 앞자리 0 무시, 정/부 칸
        fco_people = [("X999", "하람병원", "MR", "김민준"), ("C555", "(재)하람 병원", "CT/AMI", "Closed담당")]
        st = main.apply_fco_people(rows, fco_people)
        assert rows[5]["fse"] == "김민준" and rows[5]["ppsrc"] == "fco", rows[5]
        assert rows[6]["fse"] == "", rows[6]                      # "Closed담당"은 이름 형식 아님 → CT 담당 없음
        assert st.get("fco_name") == 1, st


def test_ragged_rows():
    """빠른 읽기 모드에선 행 길이가 제각각 → 'list index out of range' 오류(2026-10-31) 재발 방지"""
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        # FCO: 헤더 다음 행이 짧음
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Action Plan_MR"
        ws.append(["FCO", "Due Date", "FCO Title", "SN", "Account", "FSE", "Status", "Remark"])
        ws.append(["FCO1", dt.date(2027, 1, 1), "t", "S1", "가람병원", "김민준"])
        ws.append(["FCO2"])
        ws.append(["FCO3", None, None, None, None, None, "Closed", "x", "넘치는 칸"])
        wb.save(tmp / "fco.xlsx")
        f = main.parse_fco(tmp / "fco.xlsx", "MR", dt.date(2026, 10, 31))
        assert f["ok"] and len(f["rows"]) == 2 and f["closed"] == 1, f
        # 고객안내문: 제목 행 + 짧은 헤더 + 긴 데이터 행
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "MR"
        ws.append(["2026 고객안내문"])
        ws.append(["C&R#", "일련번호(S/N)", "국문병원명", "최종증빙 확인"])
        ws.append(["C1", "62030", "가람병원", "O", "비고", "더 긴 행"])
        ws.append(["C1", "62031", "누리병원"])
        wb.save(tmp / "fsn.xlsx")
        r = main.parse_fsn(tmp / "fsn.xlsx")
        assert r["ok"] and [x["s"] for x in r["rows"]] == [1, 0], r
        # 담당자 파일: 짧은 행/긴 행 섞임
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "FSE"
        ws.append(["FSE", "담당 병원"])
        ws.append(["김민준", "가람병원", "x", "y"])
        ws.append(["이서연"])
        ws.append(["박지훈", "누리병원"])
        wb.save(tmp / "pm.xlsx")
        pp = main.parse_people(tmp / "pm.xlsx")
        assert pp["ok"], pp
        # 파일이 깨져 있어도 전체가 멈추지 않음
        (tmp / "broken.xlsx").write_bytes(b"not an excel file")
        bad = main.safe_parse("테스트", main.parse_fco, tmp / "broken.xlsx", "MR", dt.date(2026, 10, 31))
        assert bad["ok"] is False and bad["error"]


def test_modality_from_pm_system_tabs():
    """PM Planning의 MR/CT/XR system 탭으로 장비 구분 (2026-10-31 사용자 요청)"""
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "pm.xlsx"
        wb = openpyxl.Workbook()
        f = wb.active
        f.title = "FSE"
        f.append(["FSE"])
        for nm in ("김민준", "이서연", "박지훈"):
            f.append([nm])
        for title, rows in (("MR System", [("MRS10001", "김민준"), ("MRS10002", "김민준")]),
                            ("CT_System", [("CTS20001", "이서연"), ("XRS90009", "이서연")]),
                            ("XR system", [("XRS30001", "박지훈"), ("XRS30002", "박지훈")])):
            w = wb.create_sheet(title)
            w.append(["S/N", "System", "FSE"])
            for sn, fse in rows:
                w.append([sn, "장비", fse])
        wb.save(path)
        pp = main.parse_people(path)
        eq = pp["equipment"]
        assert {x["mod"] for x in eq["sheets"]} == {"MR", "CT", "XR"}, eq["sheets"]
        assert eq["sn"]["MRS10001"][0]["mod"] == "MR" and eq["sn"]["XRS30001"][0]["mod"] == "XR"
        assert eq["sn"]["XRS90009"][0]["mod"] == "CT"

        def row(sn, bu, sys=""):
            return {"h": "가람병원", "hen": "", "bu": bu, "sn": sn, "eq": "", "code": "", "sys": sys, "fse": "", "sales": ""}
        rows = [row("XRS30001", "IGT DXR US HPM"),          # PM XR 탭 → XR
                row("XRS90009", "MR"),                      # 시트는 MR이지만 PM 장비 탭(CT)이 우선
                row("", "MR"),                        # S/N 없음 → 시트 이름 MR
                row("", "IGT DXR US HPM", "Azurion 7 M20"),   # 섞인 시트 → 장비명으로 XR
                row("", "IGT DXR US HPM", "EPIQ Elite"),      # 초음파 → 기타
                row("ZZZ99999", "IGT DXR US HPM", "")]              # 못 정함 → 기타
        res = main.assign_modality(rows, eq, [])
        assert [r["mod"] for r in rows] == ["XR", "CT", "MR", "XR", "기타", "기타"], [r["mod"] for r in rows]
        assert res["counts"]["XR"] == 2 and res["how"]["pm_sn"] == 2
        # 담당 FSE도 장비 종류에 맞게: X1(XR) → XR 탭 담당
        main.apply_people(rows, [], None, pp)
        assert rows[0]["fse"] == "박지훈" and rows[1]["fse"] == "이서연", (rows[0], rows[1])


def test_igt_sheet_korean_names():
    """IGT DXR US HPM 시트(2026-10-31 캡처): '고객명(국문)'은 영문, '고객주소 (설치장소)'가 국문 병원명"""
    import openpyxl
    head = ["C&R#", "장비 IB", "장비코드", "장비명", "일련번호(S/N)", "C&R Title", "Modality", "IB", "Customer No.",
            "Region", "한글고객명", "고객명(영문) - 삭제", "고객명(국문)", "고객주소 (설치장소)", "우편번호", "최종증빙 확인"]
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "fsn.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "IGT DXR US HPM"
        ws.append(head)
        ws.append(["2022-IGT-BST-007", "77073772", "722078", "Azurion 7 M12", "367", "oFinger", "N", "N", "N", "CC", "0",
                   "0", "Daejoen St. Mary Hosp.", "가톨릭대학교 대전성모병원", "34943", "O"])
        ws.append(["2022-IGT-BST-007", "88571195", "722078", "Azurion 7 M12", "980", "oFinger",
                   "minyoung0305@yuhs.ac", "N", "N", "Central", "0", "0", "Gangnam Severance Hospital", "강남세브란스병원",
                   "06273", ""])
        ws.append(["2022-IGT-BST-007", "67938926", "722058", "Allura Xper FD20", "115", "oFinger", "0", "0", "0",
                   "Central 2", "0", "0", "Galen Foundation(PMC Park)", "PMC박병원", "17788", ""])
        ws.append(["2022-IGT-BST-007", "39571615", "722003", "Allura Xper FD20", "1037", "oFinger", "NA", "NA", "NA",
                   "NA", "0", "0", "-", "UNKN/SCRP", "NA", ""])
        m = wb.create_sheet("MR")
        m.append(["C&R#", "장비 ID", "장비명", "일련번호(S/N)", "영문병원명(업데이트예정)", "국문병원명", "최종증빙 확인"])
        m.append(["2026-PD-MR-004", "98162875", "Ingenia", "62030", "DONGKOOK UNIV", "동국대학교병원", ""])
        m.append(["2026-PD-MR-004", "98747245", "Ingenia", "62039", "-", "", ""])
        wb.save(path)
        r = main.parse_fsn(path)
        assert r["ok"], r["error"]
        igt = [x for x in r["rows"] if x["sheet"] == "IGT DXR US HPM"]
        assert [x["h"] for x in igt] == ["가톨릭대학교 대전성모병원", "강남세브란스병원", "PMC박병원"], [x["h"] for x in igt]
        assert igt[0]["hen"] == "Daejoen St. Mary Hosp.", igt[0]
        assert all("@" not in x["bu"] for x in r["rows"]), [x["bu"] for x in r["rows"]]   # Modality(이메일)는 BU 아님
        mr = [x for x in r["rows"] if x["sheet"] == "MR"]
        assert [x["h"] for x in mr] == ["동국대학교병원"], mr     # 병원명 없는 행은 건너뜀
        res = main.assign_modality(r["rows"], {}, [])
        assert [x["mod"] for x in igt] == ["XR", "XR", "XR"] and mr[0]["mod"] == "MR", [x["mod"] for x in r["rows"]]


def test_same_hospital():
    yes = [("가톨릭대학교 인천성모병원", "인천성모 Ing CX 3.0T"), ("분당서울대학교병원", "분당서울대 31번 Ing Elition X 3.0T"),
           ("군포샘병원", "군포샘 Ach3.0T TX(효산의료재단)"), ("중앙대학교광명병원", "중앙대광명 소하검진 Ing Elition X 3.0T"),
           ("강릉동인병원", "강릉동인 Azurion 7"), ("(재)가람 의료원", "가람의료원")]
    no = [("강릉동인병원", "강릉아산병원 Azurion"), ("강릉동인병원", "다른병원 Allura"), ("서울대학교병원", "서울성모 Ing")]
    for a_, b_ in yes:
        assert main.same_hospital([a_], [b_]), (a_, b_)
    for a_, b_ in no:
        assert not main.same_hospital([a_], [b_]), (a_, b_)


def test_short_serial_collision():
    """짧은 S/N이 다른 병원 장비와 겹쳐 엉뚱한 FSE가 붙던 문제(2026-10-31 강릉동인병원에 장재훈) 재발 방지"""
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "pm.xlsx"
        wb = openpyxl.Workbook()
        f = wb.active
        f.title = "FSE"
        f.append(["FSE"])
        for nm in ("이상우", "장재훈", "김민준"):
            f.append([nm])
        x = wb.create_sheet("XR System")
        x.append(["S/N", "Site", "System", "FSE"])
        x.append(["370", "강릉동인 Azurion 7", "Azurion 7 M12", "이상우"])
        x.append(["370", "다른병원 Allura", "Allura Xper FD20", "장재훈"])   # 같은 번호, 다른 병원
        x.append(["371", "다른병원 Allura", "Allura Xper FD20", "장재훈"])
        m = wb.create_sheet("MR System")
        m.append(["S/N", "Site", "System", "FSE"])
        m.append(["370", "강릉동인 Ingenia", "Ingenia 1.5T", "김민준"])         # 같은 번호, 다른 종류(MR)
        m.append(["MR778899", "아무병원", "Ingenia 3.0T", "김민준"])
        wb.save(path)
        pp = main.parse_people(path)
        assert pp["ok"], pp

        def row(h, sn, sheet, sys=""):
            return {"h": h, "hen": "", "bu": "", "sheet": sheet, "sn": sn, "eq": "", "code": "", "sys": sys,
                    "fse": "", "sales": ""}
        rows = [row("강릉동인병원", "370", "IGT DXR US HPM", "Azurion 7 M12"),   # XR · 강릉동인 → 이상우만
                row("새병원", "371", "IGT DXR US HPM", "Allura Xper FD20"),      # 번호는 같지만 병원이 다름 → 연결 안 함
                row("강릉동인병원", "370", "MR", "Ingenia 1.5T"),                 # MR · 강릉동인 → 김민준
                row("이름다른병원", "MR778899", "MR")]                            # 긴 S/N은 병원명이 달라도 인정
        main.assign_modality(rows, pp["equipment"], [])
        assert [r["mod"] for r in rows] == ["XR", "XR", "MR", "MR"], [r["mod"] for r in rows]
        main.apply_people(rows, [], None, pp)
        assert rows[0]["fse"] == "이상우", rows[0]
        assert rows[0]["ppwhy"].startswith("담당자 파일 'XR System'"), rows[0]
        assert rows[1]["fse"] == "", rows[1]
        assert rows[2]["fse"] == "김민준", rows[2]
        assert rows[3]["fse"] == "김민준", rows[3]
        # FCO 보충도 같은 기준: 짧은 S/N + 다른 병원 → 연결 안 함
        r5 = row("새병원", "371", "IGT DXR US HPM")
        r5["mod"] = "XR"
        main.apply_fco_people([r5], [("371", "다른병원", "XR", "장재훈")])
        assert r5["fse"] == "", r5


def test_download_urls():
    fco = main.FCO_SOURCES["MR"]["url"]
    assert main.add_download_param(fco) == (
        "https://share.philips.com/sites/KRHSFCOstatus/_layouts/15/download.aspx"
        "?UniqueId=41877D54-C06C-43EA-A61A-D8EE1F661BF5"), main.add_download_param(fco)
    assert main.add_download_param(main.FCO_SOURCES["CT/AMI"]["url"]).endswith(
        "UniqueId=FAACFA92-78CB-4165-BEEB-C5346C578244")
    fsn = main.FSN_SOURCE["url"]
    assert main.add_download_param(fsn) == fsn + "&download=1"
    assert all(v["url"].startswith("https://") for v in main.FCO_SOURCES.values())


def test_pipeline_manual_file():
    """파일 직접 선택 경로: 다운로드 없이 대시보드 생성"""
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        _samples(tmp)
        old = (main.APP_DIR, main.DATA_DIR, main.DOWNLOAD_DIR, main.LOG_PATH, main.DASHBOARD_PATH,
               main.CONFIG_PATH, main.FCO_APP_DIR)
        try:
            main.APP_DIR = tmp / "app"
            main.DATA_DIR = main.APP_DIR / "data"
            main.DOWNLOAD_DIR = main.APP_DIR / "downloads"
            main.LOG_PATH = main.APP_DIR / "debug_log.txt"
            main.DASHBOARD_PATH = main.APP_DIR / "dash.html"
            main.CONFIG_PATH = main.APP_DIR / "config.json"
            main.FCO_APP_DIR = tmp / "fco_app"
            saved_urls = {k: v["url"] for k, v in main.FCO_SOURCES.items()}
            for v in main.FCO_SOURCES.values():
                v["url"] = ""  # 테스트에서는 Edge로 받지 않음
            saved_pm = main.PM_SOURCE["url"]
            main.PM_SOURCE["url"] = ""
            main.FCO_APP_DIR.mkdir()
            (tmp / "MR FCO Action Plan.xlsx").replace(main.FCO_APP_DIR / "MR_FCO_Action_Plan.xlsx")
            steps = []
            out = main.run_pipeline(False, tmp / "A.xlsx", lambda t: None, lambda i, s, t: steps.append((i, s)),
                                    threading.Event(), {})
            assert out.exists()
            assert (0, "run") in steps and (3, "ok") in steps, steps
            assert (2, "warn") in steps, steps  # 담당자 파일 없음 → 경고만
            assert (main.DATA_DIR / main.FSN_SOURCE["cache"]).exists()
        finally:
            if "saved_pm" in dir():
                main.PM_SOURCE["url"] = saved_pm
            for k, v in main.FCO_SOURCES.items():
                v["url"] = saved_urls.get(k, v["url"]) if "saved_urls" in dir() else v["url"]
            (main.APP_DIR, main.DATA_DIR, main.DOWNLOAD_DIR, main.LOG_PATH, main.DASHBOARD_PATH,
             main.CONFIG_PATH, main.FCO_APP_DIR) = old


if __name__ == "__main__":
    failed = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS  {name}")
            except Exception as e:  # noqa: BLE001
                failed += 1
                print(f"FAIL  {name}: {e}")
    print("모두 통과" if not failed else f"{failed}개 실패")
    sys.exit(1 if failed else 0)
