# -*- coding: utf-8 -*-
"""
개발/테스트용 샘플 엑셀 생성 (실제 데이터 아님, 병원·사람 이름은 모두 가상).

python tools/make_sample_data.py <출력폴더>
  → FSN_sample_A.xlsx  (일반 표: 전달여부/전달방법 컬럼)
  → FSN_sample_B.xlsx  (두 줄 헤더 + 병합 셀 + 등기/직접전달/이메일 칸 분리, 전달여부 컬럼 없음)
  → MR FCO Action Plan.xlsx / XR Open FCO list.xlsx / CT_AMI FCO Action Plan.xlsx
"""
import datetime as dt
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

random.seed(7)
TODAY = dt.date(2026, 10, 23)

HOSPITALS = [
    "한빛대학교병원", "가람의료원", "누리병원", "새솔병원", "다온병원", "미르병원", "보람의료재단 보람병원",
    "아라병원", "이음병원", "채움병원", "하람병원", "다솜병원", "늘봄병원", "소망병원", "푸른숲병원",
    "도담병원", "온새미로병원", "별빛병원", "해오름병원", "가온대학교병원", "나래병원", "라온병원",
]
FSE = ["김민준", "이서연", "박지훈", "최유진", "정하늘", "강도윤"]
SALES = ["윤서준", "장미래", "임하준", "한지우"]
BUS = ["MR", "CT", "DXR", "IGT", "US"]
SYSTEMS = {"MR": ["Ingenia 1.5T", "Ingenia Elition 3.0T", "Achieva 1.5T"],
           "CT": ["Incisive CT", "IQon Spectral CT", "Brilliance 64"],
           "DXR": ["DigitalDiagnost C90", "MobileDiagnost wDR"],
           "IGT": ["Azurion 7 M20", "Allura Xper FD20"],
           "US": ["EPIQ Elite", "Affiniti 70"]}
FSNS = [
    ("2026-FSN-001", "C&R 2026-0112", "MR", "코일 커넥터 점검 안내"),
    ("2026-FSN-002", "C&R 2026-0135", "CT", "소프트웨어 업데이트 안내"),
    ("2026-FSN-003", "C&R 2026-0151", "IGT", "테이블 브레이크 점검 안내"),
    ("2026-FSN-004", "C&R 2026-0178", "DXR", "디텍터 배터리 교체 안내"),
    ("2026-FSN-005", "C&R 2026-0190", "US", "프로브 표면 점검 안내"),
    ("2026-FSN-006", "C&R 2026-0203", "MR", "헬륨 모니터링 안내"),
]
HOSP_FSE = {h: random.choice(FSE) for h in HOSPITALS}
HOSP_SALES = {h: random.choice(SALES) for h in HOSPITALS}


def make_equipment():
    eq = []
    for h in HOSPITALS:
        for bu in random.sample(BUS, random.randint(1, 4)):
            for _ in range(random.randint(1, 2)):
                eq.append({"h": h, "bu": bu, "sys": random.choice(SYSTEMS[bu]),
                           "sn": f"{bu[:2]}{random.randint(10000, 99999)}"})
    return eq


EQUIP = make_equipment()


def fsn_rows():
    rows = []
    for fsn, cr, bu, title in FSNS:
        targets = [e for e in EQUIP if e["bu"] == bu]
        for e in targets:
            r = random.random()
            if r < 0.62:
                status = random.choice(["완료", "전달완료", "O", "완료"])
                method = random.choices(["등기", "직접전달", "이메일", "등기, 이메일"], [5, 3, 3, 1])[0]
                date = TODAY - dt.timedelta(days=random.randint(3, 120))
            elif r < 0.95:
                status = random.choice(["미전달", "", "진행중", "X"])
                method = random.choice(["", "", "등기", "직접전달"])
                date = None
            else:
                status = "해당없음"
                method = ""
                date = None
            rows.append({"fsn": fsn, "cr": cr, "bu": bu, "title": title, **e, "status": status, "method": method,
                         "date": date, "fse": HOSP_FSE[e["h"]], "sales": HOSP_SALES[e["h"]],
                         "reply": "회신완료" if (date and random.random() < 0.6) else "",
                         "remark": random.choice(["", "", "", "담당자 부재로 재방문 예정", "원내 메일 전달"])})
    return rows


def sample_a(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "전달현황"
    ws["A1"] = "2026 고객안내문(FSN) 전달 현황"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:O1")
    head = ["No", "BU", "C&R", "FSN 번호", "제목", "병원명", "S/N", "장비", "담당 FSE", "영업담당", "전달방법", "전달여부",
            "전달일", "회신여부", "비고"]
    ws.append([])
    ws.append(head)
    for i, r in enumerate(fsn_rows(), 1):
        ws.append([i, r["bu"], r["cr"], r["fsn"], r["title"], r["h"], r["sn"], r["sys"], r["fse"], r["sales"],
                   r["method"], r["status"], r["date"], r["reply"], r["remark"]])
    for row in ws.iter_rows(min_row=4, min_col=13, max_col=13):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    # 요약 시트 (무시되어야 함)
    s = wb.create_sheet("요약")
    s.append(["BU", "대상", "완료", "미전달"])
    for bu in BUS:
        s.append([bu, 10, 7, 3])
    wb.save(path)


def sample_b(path: Path):
    """두 줄 헤더, 병원명 세로 병합, 전달방법이 칸별로 나뉨, 전달여부 컬럼 없음"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FSN_2026-0151"
    ws.append(["병원", "Modality", "FSN", "전달방법", None, None, "담당자", "비고"])
    ws.append([None, None, None, "등기", "직접전달", "이메일", None, None])
    ws.merge_cells("D1:F1")
    for col in "ABCGH":
        ws.merge_cells(f"{col}1:{col}2")
    rows = [r for r in fsn_rows() if r["fsn"] in ("2026-FSN-003", "2026-FSN-001")]
    rows.sort(key=lambda r: r["h"])
    start = 3
    i = 0
    while i < len(rows):
        h = rows[i]["h"]
        j = i
        while j < len(rows) and rows[j]["h"] == h:
            r = rows[j]
            done = r["status"] in ("완료", "전달완료", "O")
            marks = {"등기": None, "직접전달": None, "이메일": None}
            if done:
                for m in r["method"].replace(" ", "").split(","):
                    if m in marks:
                        marks[m] = r["date"] if m == "등기" else "O"
            ws.append([h if j == i else None, r["bu"], r["fsn"], marks["등기"], marks["직접전달"], marks["이메일"],
                       r["fse"], r["remark"]])
            j += 1
        if j - i > 1:
            ws.merge_cells(f"A{start}:A{start + (j - i) - 1}")
        start += j - i
        i = j
    for c in ws["A"]:
        c.alignment = Alignment(vertical="center")
    wb.save(path)


FCO_HEAD = ["FCO", "Due Date", "FCO Title", "IP", "System Code", "System Code Description", "SN", "Account", "Site",
            "Case", "WO", "Region", "FSE", "SW Ver", "Action Plan", "Status", "고객안내문 관련", "Remark", "PM Cycle",
            "Kit 1"]


def fco_file(path: Path, src: str, bus):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Action Plan_{src.split('/')[0]}"
    ws.append(FCO_HEAD)
    targets = [e for e in EQUIP if e["bu"] in bus]
    n = 0
    for k in range(6):
        fco = f"FCO{random.randint(70000000, 79999999)}"
        due = TODAY + dt.timedelta(days=random.choice([-20, 40, 100, 150, 250, 300, 420, 600]))
        title = random.choice(["Gradient coil inspection", "SW update R11.2", "Tube cable replacement",
                               "Table motor check", "Detector firmware update", "Cooling unit inspection"])
        rel = random.choice(["", "", "대상 (2026-FSN-00%d)" % random.randint(1, 6)])
        for e in random.sample(targets, min(len(targets), random.randint(4, 12))):
            status = random.choices(["Open", "In Progress", "Closed", "Scheduled"], [5, 3, 3, 2])[0]
            # 일부 병원명은 표기를 다르게 (S/N 연결 테스트)
            account = e["h"] if random.random() < 0.8 else "(재)" + e["h"].replace("병원", " 병원")
            ws.append([fco, due, title, "", e["bu"] + "100", e["sys"], e["sn"], account, "본관", f"C{random.randint(1000, 9999)}",
                       f"WO{random.randint(100000, 999999)}", random.choice(["서울", "경기", "부산", "대구"]),
                       HOSP_FSE[e["h"]], "", random.choice(["부품 입고 대기", "고객 일정 협의", "PM 시 진행", ""]),
                       status, rel, "", "6M", ""])
            n += 1
    wb.save(path)
    return n


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "sample_data")
    out.mkdir(parents=True, exist_ok=True)
    sample_a(out / "FSN_sample_A.xlsx")
    sample_b(out / "FSN_sample_B.xlsx")
    fco_file(out / "MR FCO Action Plan.xlsx", "MR", ["MR"])
    fco_file(out / "XR Open FCO list.xlsx", "XR", ["DXR", "IGT"])
    fco_file(out / "CT_AMI FCO Action Plan.xlsx", "CT/AMI", ["CT"])
    print("샘플 생성:", out.resolve())


if __name__ == "__main__":
    main()
